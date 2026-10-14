"""Tables: header and value classification, apply to xlsx and csv.

Detectors and the replacer are swapped for small fakes, so these tests do not
need natasha or spaCy. All data is invented.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from anon import table
from anon.model import Kind, Merged
from anon.vault import Vault, fold

# ---- fakes ---------------------------------------------------------------

# (surface, normal) pairs the fake NER knows.
PEOPLE = [
    ("Ковалева Мария Ивановна", "Ковалева Мария Ивановна"),
    ("Петров Иван Сергеевич", "Петров Иван Сергеевич"),
    ("Сидорова Анна Павловна", "Сидорова Анна Павловна"),
    ("Мария Ковалева", "Мария Ковалева"),
    ("Марией Ковалевой", "Мария Ковалева"),
    ("Ивана Петрова", "Иван Петров"),
    ("Глеб Орлов", "Глеб Орлов"),
    ("Олег Зайцев", "Олег Зайцев"),
    ("John Smith", "John Smith"),
    ("Emily Clark", "Emily Clark"),
    ("Mark Brown", "Mark Brown"),
]
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
_PHONE = re.compile(r"\+\d[\d\s()-]{8,}\d")


def fake_find_all(text, langs=None, layers=None):
    out: list[Merged] = []
    taken = [False] * len(text)
    for surface, normal in sorted(PEOPLE, key=lambda p: -len(p[0])):
        for m in re.finditer(re.escape(surface), text):
            if any(taken[m.start() : m.end()]):
                continue
            taken[m.start() : m.end()] = [True] * (m.end() - m.start())
            out.append(
                Merged(m.start(), m.end(), m.group(0), Kind.PERSON, ["natasha"], normal)
            )
    for rx, kind in ((_EMAIL, Kind.EMAIL), (_PHONE, Kind.PHONE)):
        for m in rx.finditer(text):
            if not any(taken[m.start() : m.end()]):
                out.append(Merged(m.start(), m.end(), m.group(0), kind, ["rules"]))
    return sorted(out, key=lambda s: s.start)


def fake_apply(text, spans, vault, lang):
    for s in sorted(spans, key=lambda s: s.start, reverse=True):
        if vault.is_ignored(s.text):
            continue
        if s.kind == Kind.PERSON:
            e = table.person_entity(vault, s.normal or s.text)
            e.variants.add(fold(s.text))
        else:
            e = vault.get_or_create(s.kind, s.text, s.normal)
        text = text[: s.start] + vault.label(e, lang) + text[s.end :]
    return text


@pytest.fixture(autouse=True)
def fakes(monkeypatch):
    monkeypatch.setattr(table, "_detector", lambda: fake_find_all)
    monkeypatch.setattr(table, "_replacer", lambda: fake_apply)


LABEL = re.compile(r"^\[[A-ZА-Я]+_\d+\]$")


def _values(path: Path, sheet: str | None = None) -> list[list]:
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _all_text(path: Path) -> str:
    wb = load_workbook(path)
    return "\n".join(
        " ".join("" if v is None else str(v) for v in row)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
    )


# ---- header classification ----------------------------------------------


@pytest.mark.parametrize(
    "header, kind, part",
    [
        ("ФИО", Kind.PERSON, None),
        ("Ф.И.О.", Kind.PERSON, None),
        ("Сотрудник", Kind.PERSON, None),
        ("Руководитель отдела", Kind.PERSON, None),
        ("Ответственный", Kind.PERSON, None),
        ("Full name", Kind.PERSON, None),
        ("Name", Kind.PERSON, None),
        ("Department manager", Kind.PERSON, None),
        ("Owner", Kind.PERSON, None),
        ("Фамилия", Kind.PERSON, "last"),
        ("Имя", Kind.PERSON, "first"),
        ("Отчество", Kind.PERSON, "middle"),
        ("Last name", Kind.PERSON, "last"),
        ("First Name", Kind.PERSON, "first"),
        ("Patronymic", Kind.PERSON, "middle"),
        ("E-mail", Kind.EMAIL, None),
        ("Эл. почта", Kind.EMAIL, None),
        ("Email address", Kind.EMAIL, None),
        ("Тел.", Kind.PHONE, None),
        ("Mobile", Kind.PHONE, None),
        ("Телефон руководителя", Kind.PHONE, None),
        ("СНИЛС", Kind.SNILS, None),
        ("ИНН", Kind.INN, None),
        ("Паспортные данные", Kind.PASSPORT, None),
        ("SSN", Kind.SSN, None),
        ("NI number", Kind.NINO, None),
        ("Табельный номер", Kind.ID, None),
        ("Таб. №", Kind.ID, None),
        ("Employee ID", Kind.ID, None),
        ("Personnel number", Kind.ID, None),
        ("Адрес", Kind.LOCATION, None),
        ("Почтовый адрес", Kind.LOCATION, None),
        ("Home city", Kind.LOCATION, None),
        ("Город проживания", Kind.LOCATION, None),
        # traps
        ("Наименование должности", None, None),
        ("Название отдела", None, None),
        ("Имя файла", None, None),
        ("Отдел сотрудника", None, None),
        ("Должность", None, None),
        ("Количество сотрудников", None, None),
        ("Project name", None, None),
        ("File name", None, None),
        ("Employee department", None, None),
        ("Оклад", None, None),
        ("Комментарий", None, None),
        ("Department", None, None),
    ],
)
def test_classify_header(header, kind, part):
    assert table.classify_header(header) == (kind, part)


# ---- ru HR table ---------------------------------------------------------


def _ru_hr(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Штат"
    ws.append(
        [
            "ФИО",
            "Табельный номер",
            "Наименование должности",
            "Оклад",
            "Отдел",
            "Годовой",
            "Комментарий",
            "Дата приема",
        ]
    )
    rows = [
        (
            "Ковалева Мария Ивановна",
            1001,
            "Аналитик",
            120000,
            "Финансы",
            "Согласовано с Ивана Петрова",
            date(2023, 3, 1),
        ),
        (
            "Петров Иван Сергеевич",
            1002,
            "Руководитель",
            250000,
            "Финансы",
            "Замещает Марией Ковалевой в отпуске",
            date(2021, 9, 15),
        ),
        (
            "Сидорова Анна Павловна",
            1003,
            "Юрист",
            150000,
            "Правовой",
            "",
            date(2022, 1, 10),
        ),
    ]
    for i, (fio, tab, pos, pay, dep, note, hired) in enumerate(rows, start=2):
        ws.append([fio, tab, pos, pay, dep, f"=D{i}*12", note or None, hired])
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 33
    path = path / "hr.xlsx"
    wb.save(path)
    return path


def test_ru_hr_scan(tmp_path):
    scan = table.scan_table(_ru_hr(tmp_path))
    assert scan.format == "xlsx"
    sh = scan.sheets[0]
    assert sh.header_row == 1 and sh.data_start == 2
    cols = {c.header: c for c in sh.columns.values()}
    assert (cols["ФИО"].kind, cols["ФИО"].reason) == (Kind.PERSON, "header")
    assert (cols["Табельный номер"].kind, cols["Табельный номер"].reason) == (
        Kind.ID,
        "header",
    )
    for h in (
        "Наименование должности",
        "Оклад",
        "Отдел",
        "Годовой",
        "Комментарий",
        "Дата приема",
    ):
        assert cols[h].kind is None, h
    assert cols["Оклад"].numeric and cols["Дата приема"].dates
    # The comment column is free text: its names are review items with context.
    notes = [it for it in scan.items() if it.cell.startswith("G")]
    assert {it.text for it in notes} == {"Ивана Петрова", "Марией Ковалевой"}
    assert all(
        it.context.startswith("Комментарий") and it.sheet == "Штат" for it in notes
    )
    fio = [it for it in scan.items() if it.cell == "A2"][0]
    assert fio.kind == Kind.PERSON and fio.text == "Ковалева Мария Ивановна"
    assert fio.sources == ["column"] and "ФИО" in fio.context


def test_ru_hr_apply(tmp_path):
    src = _ru_hr(tmp_path)
    scan = table.scan_table(src)
    vault = Vault()
    out = tmp_path / "out.xlsx"
    table.apply_table(src, scan, vault, out)

    rows = _values(out)
    assert rows[0] == [
        "ФИО",
        "Табельный номер",
        "Наименование должности",
        "Оклад",
        "Отдел",
        "Годовой",
        "Комментарий",
        "Дата приема",
    ]
    labels = [r[0] for r in rows[1:]]
    assert labels == ["[ИМЯ_1]", "[ИМЯ_2]", "[ИМЯ_3]"]
    assert [r[1] for r in rows[1:]] == ["[НОМЕР_1]", "[НОМЕР_2]", "[НОМЕР_3]"]
    assert [r[2] for r in rows[1:]] == ["Аналитик", "Руководитель", "Юрист"]
    assert [r[3] for r in rows[1:]] == [120000, 250000, 150000]
    assert [r[5] for r in rows[1:]] == ["=D2*12", "=D3*12", "=D4*12"]
    assert rows[1][7].date() == date(2023, 3, 1)
    # Names in comments resolve to the people of the ФИО column.
    assert rows[1][6] == "Согласовано с [ИМЯ_2]"
    assert rows[2][6] == "Замещает [ИМЯ_1] в отпуске"

    text = _all_text(out)
    for name in ("Ковалева", "Петров", "Сидорова", "Ковалевой", "Петрова"):
        assert name not in text

    wb = load_workbook(out)
    ws = wb["Штат"]
    assert ws["A1"].font.bold
    assert ws.column_dimensions["A"].width == 33
    # The source file is untouched.
    assert _values(src)[1][0] == "Ковалева Мария Ивановна"


# ---- en staff table ------------------------------------------------------


def test_en_staff(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(["Full name", "Email", "Phone", "Department", "Salary"])
    ws.append(
        ["John Smith", "john.smith@example.com", "+44 20 7946 0018", "Sales", 52000]
    )
    ws.append(
        ["Emily Clark", "e.clark@example.com", "+44 20 7946 0019", "Finance", 61000]
    )
    ws.append(
        ["John Smith", "john.smith@example.com", "+44 20 7946 0018", "Sales", 52000]
    )
    src = tmp_path / "staff.xlsx"
    wb.save(src)

    scan = table.scan_table(src)
    kinds = {c.header: c.kind for c in scan.sheets[0].columns.values()}
    assert kinds == {
        "Full name": Kind.PERSON,
        "Email": Kind.EMAIL,
        "Phone": Kind.PHONE,
        "Department": None,
        "Salary": None,
    }
    vault = Vault()
    out = tmp_path / "staff.anon.xlsx"
    table.apply_table(src, scan, vault, out)  # lang auto: en headers
    rows = _values(out)
    assert rows[1][:3] == ["[PERSON_1]", "[EMAIL_1]", "[PHONE_1]"]
    assert rows[2][:3] == ["[PERSON_2]", "[EMAIL_2]", "[PHONE_2]"]
    assert rows[3][:3] == rows[1][:3]  # same person, same labels
    assert [r[3:] for r in rows[1:]] == [
        ["Sales", 52000],
        ["Finance", 61000],
        ["Sales", 52000],
    ]
    assert "Smith" not in _all_text(out) and "example.com" not in _all_text(out)


# ---- split name columns + same person across sheets ----------------------


def test_split_name_columns_and_cross_sheet_identity(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Состав"
    ws.append(["Фамилия", "Имя", "Отчество", "Должность"])
    ws.append(["Ковалева", "Мария", "Ивановна", "Аналитик"])
    ws.append(["Петров", "Иван", "Сергеевич", "Руководитель"])
    ws2 = wb.create_sheet("Проекты")
    ws2.append(["Проект", "Ответственный", "Бюджет"])
    ws2.append(["Аудит", "Мария Ковалева", 300000])
    ws2.append(["Отчет", "Петров Иван Сергеевич", 120000])
    src = tmp_path / "split.xlsx"
    wb.save(src)

    scan = table.scan_table(src)
    parts = {c.header: (c.kind, c.part) for c in scan.sheet("Состав").columns.values()}
    assert parts["Фамилия"] == (Kind.PERSON, "last")
    assert parts["Имя"] == (Kind.PERSON, "first")
    assert parts["Отчество"] == (Kind.PERSON, "middle")
    assert parts["Должность"] == (None, None)
    person_items = [it for it in scan.items() if it.sheet == "Состав"]
    assert [it.text for it in person_items] == [
        "Ковалева Мария Ивановна",
        "Петров Иван Сергеевич",
    ]
    assert person_items[0].cell == "A2"

    vault = Vault()
    out = tmp_path / "split.anon.xlsx"
    table.apply_table(src, scan, vault, out, lang="ru")
    rows = _values(out, "Состав")
    assert rows[1] == ["[ИМЯ_1]", "[ИМЯ_1]", "[ИМЯ_1]", "Аналитик"]
    assert rows[2] == ["[ИМЯ_2]", "[ИМЯ_2]", "[ИМЯ_2]", "Руководитель"]
    rows2 = _values(out, "Проекты")
    assert [r[1] for r in rows2[1:]] == ["[ИМЯ_1]", "[ИМЯ_2]"]
    assert [r[2] for r in rows2[1:]] == [300000, 120000]
    assert len([e for e in vault.entities if e.kind == Kind.PERSON]) == 2
    assert "Ковалева" not in _all_text(out) and "Мария" not in _all_text(out)


def test_same_person_two_sheets_and_two_files_share_vault(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Январь"
    ws.append(["Сотрудник", "Часы"])
    ws.append(["Глеб Орлов", 160])
    ws2 = wb.create_sheet("Февраль")
    ws2.append(["Сотрудник", "Часы"])
    ws2.append(["Олег Зайцев", 150])
    ws2.append(["Глеб Орлов", 152])
    src = tmp_path / "hours.xlsx"
    wb.save(src)
    other = tmp_path / "bonus.csv"
    other.write_text("ФИО,Премия\nГлеб Орлов,10000\n", encoding="utf-8")

    vault = Vault()
    table.apply_table(src, table.scan_table(src), vault, tmp_path / "hours.anon.xlsx")
    table.apply_table(
        other, table.scan_table(other), vault, tmp_path / "bonus.anon.csv"
    )
    jan = _values(tmp_path / "hours.anon.xlsx", "Январь")
    feb = _values(tmp_path / "hours.anon.xlsx", "Февраль")
    assert jan[1][0] == feb[2][0] == "[ИМЯ_1]"
    assert feb[1][0] == "[ИМЯ_2]"
    bonus = (tmp_path / "bonus.anon.csv").read_text(encoding="utf-8-sig").splitlines()
    assert bonus == ["ФИО,Премия", "[ИМЯ_1],10000"]


# ---- csv -----------------------------------------------------------------


def test_csv_semicolon_cp1251(tmp_path):
    content = (
        "ФИО;Отдел;Оклад;Телефон;Примечание\r\n"
        "Ковалева Мария Ивановна;Финансы;120 000;+7 916 123-45-67;Подменяет Ивана Петрова\r\n"
        'Петров Иван Сергеевич;"Продажи; опт";95 000,50;+7 916 765-43-21;\r\n'
    )
    src = tmp_path / "staff.csv"
    src.write_bytes(content.encode("cp1251"))

    scan = table.scan_table(src)
    assert (scan.format, scan.delimiter, scan.encoding) == ("csv", ";", "cp1251")
    sh = scan.sheets[0]
    kinds = {c.header: c.kind for c in sh.columns.values()}
    assert kinds["ФИО"] == Kind.PERSON and kinds["Телефон"] == Kind.PHONE
    assert kinds["Отдел"] is None and kinds["Оклад"] is None
    assert sh.columns[3].numeric

    out = tmp_path / "staff.anon.csv"
    table.apply_table(src, scan, Vault(), out)
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["ФИО", "Отдел", "Оклад", "Телефон", "Примечание"]
    assert rows[1] == [
        "[ИМЯ_1]",
        "Финансы",
        "120 000",
        "[ТЕЛЕФОН_1]",
        "Подменяет [ИМЯ_2]",
    ]
    assert rows[2] == ["[ИМЯ_2]", "Продажи; опт", "95 000,50", "[ТЕЛЕФОН_2]", ""]


@pytest.mark.parametrize("delim", [",", "\t"])
def test_csv_sniffs_comma_and_tab(tmp_path, delim):
    src = tmp_path / "t.csv"
    lines = [
        ["Name", "Team", "Notes"],
        ["Mark Brown", "Ops", "ok, fine"],
        ["Emily Clark", "Ops", ""],
    ]
    buf = io.StringIO()
    csv.writer(buf, delimiter=delim).writerows(lines)
    src.write_text("﻿" + buf.getvalue(), encoding="utf-8")
    scan = table.scan_table(src)
    assert scan.delimiter == delim and scan.encoding == "utf-8-sig"
    out = tmp_path / "o.csv"
    table.apply_table(src, scan, Vault(), out)
    rows = list(
        csv.reader(io.StringIO(out.read_text(encoding="utf-8-sig")), delimiter=delim)
    )
    assert rows[1] == ["[PERSON_1]", "Ops", "ok, fine"]


# ---- value-based classification -----------------------------------------


def test_value_based_classification(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Участник", "Контактные данные", "Заметки", "Баллы"])
    ws.append(["Глеб Орлов", "gleb.orlov@example.org", "Сильный кандидат", 87])
    ws.append(["Олег Зайцев", "oleg.z@example.org", "Рекомендовал Мария Ковалева", 75])
    ws.append(["Мария Ковалева", "maria.k@example.org", "Повторное интервью", 91])
    ws.append(["вне списка", "нет", "", 60])
    src = tmp_path / "v.xlsx"
    wb.save(src)

    scan = table.scan_table(src)
    cols = scan.sheets[0].columns
    assert (cols[1].kind, cols[1].reason) == (Kind.PERSON, "values")  # 3 of 4 cells
    assert (cols[2].kind, cols[2].reason) == (Kind.EMAIL, "values")
    assert cols[3].kind is None and cols[4].kind is None
    assert [it.text for it in scan.items() if it.cell.startswith("C")] == [
        "Мария Ковалева"
    ]

    out = tmp_path / "v.anon.xlsx"
    table.apply_table(src, scan, Vault(), out)
    rows = _values(out)
    assert [r[0] for r in rows[1:4]] == ["[ИМЯ_1]", "[ИМЯ_2]", "[ИМЯ_3]"]
    assert rows[2][2] == "Рекомендовал [ИМЯ_3]"
    assert rows[4][1] == "нет"  # placeholder stays
    assert [r[3] for r in rows[1:]] == [87, 75, 91, 60]


def test_value_classification_threshold(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Колонка", "Другое"])
    for v in ["Глеб Орлов", "текст без имени", "еще текст", "и еще"]:
        ws.append([v, "x"])
    src = tmp_path / "t.xlsx"
    wb.save(src)
    scan = table.scan_table(src)
    assert scan.sheets[0].columns[1].kind is None  # 1 of 4 < 60%
    assert [it.text for it in scan.items()] == [
        "Глеб Орлов"
    ]  # still a free-text finding


# ---- header traps, merged headers, placeholders, ignored ------------------


def test_header_traps_stay(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Наименование должности", "Название отдела", "Имя файла", "Сотрудник"])
    ws.append(
        ["Ведущий аналитик", "Отдел планирования", "otchet_q1.xlsx", "Глеб Орлов"]
    )
    src = tmp_path / "traps.xlsx"
    wb.save(src)
    scan = table.scan_table(src)
    assert [c.kind for c in scan.sheets[0].columns.values()] == [
        None,
        None,
        None,
        Kind.PERSON,
    ]
    out = tmp_path / "traps.anon.xlsx"
    table.apply_table(src, scan, Vault(), out)
    assert _values(out)[1] == [
        "Ведущий аналитик",
        "Отдел планирования",
        "otchet_q1.xlsx",
        "[ИМЯ_1]",
    ]


def test_title_row_and_two_level_merged_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Сводка по персоналу за квартал"
    ws.merge_cells("A1:E1")
    ws["A2"] = "Сотрудник"
    ws.merge_cells("A2:B2")
    ws["C2"] = "Контакты"
    ws.merge_cells("C2:D2")
    ws["E2"] = "Отдел"
    ws.merge_cells("E2:E3")
    ws.append(["Фамилия", "Имя", "Email", "Телефон"])  # row 3
    ws.append(["Орлов", "Глеб", "gleb@example.org", "+7 916 000-11-22", "ИТ"])
    src = tmp_path / "merged.xlsx"
    wb.save(src)

    scan = table.scan_table(src)
    sh = scan.sheets[0]
    assert sh.header_row == 2 and sh.header_rows == [2, 3] and sh.data_start == 4
    got = {c.letter: (c.kind, c.part) for c in sh.columns.values()}
    assert got == {
        "A": (Kind.PERSON, "last"),
        "B": (Kind.PERSON, "first"),
        "C": (Kind.EMAIL, None),
        "D": (Kind.PHONE, None),
        "E": (None, None),
    }
    assert sh.columns[1].header == "Сотрудник / Фамилия"

    out = tmp_path / "merged.anon.xlsx"
    table.apply_table(src, scan, Vault(), out)
    rows = _values(out)
    assert rows[0][0] == "Сводка по персоналу за квартал"
    assert rows[2][:4] == ["Фамилия", "Имя", "Email", "Телефон"]
    assert rows[3] == ["[ИМЯ_1]", "[ИМЯ_1]", "[ПОЧТА_1]", "[ТЕЛЕФОН_1]", "ИТ"]
    assert {str(r) for r in load_workbook(out).active.merged_cells.ranges} == {
        "A1:E1",
        "A2:B2",
        "C2:D2",
        "E2:E3",
    }


def test_ignored_values_and_numeric_id_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Таб. №", "ФИО", "Дата рождения"])
    ws.append([501, "Глеб Орлов", date(1990, 5, 4)])
    ws.append([502, "Отдел кадров", date(1985, 2, 1)])
    src = tmp_path / "ids.xlsx"
    wb.save(src)
    vault = Vault()
    vault.ignored.add(fold("Отдел кадров"))
    out = tmp_path / "ids.anon.xlsx"
    table.apply_table(src, table.scan_table(src), vault, out)
    rows = _values(out)
    assert rows[1][:2] == ["[НОМЕР_1]", "[ИМЯ_1]"]
    assert rows[2][:2] == ["[НОМЕР_2]", "Отдел кадров"]
    assert rows[1][2].date() == date(1990, 5, 4)


def test_date_column_never_classified(tmp_path):
    src = tmp_path / "d.csv"
    src.write_text(
        "Сотрудник;Дата\nГлеб Орлов;01.02.2024\nОлег Зайцев;2024-03-05\n",
        encoding="utf-8",
    )
    scan = table.scan_table(src)
    assert scan.sheets[0].columns[2].dates and scan.sheets[0].columns[2].kind is None
    out = tmp_path / "d.anon.csv"
    table.apply_table(src, scan, Vault(), out)
    assert out.read_text(encoding="utf-8-sig").splitlines()[1:] == [
        "[ИМЯ_1];01.02.2024",
        "[ИМЯ_2];2024-03-05",
    ]


def test_unsupported_format(tmp_path):
    p = tmp_path / "x.ods"
    p.write_bytes(b"")
    with pytest.raises(ValueError):
        table.scan_table(p)


def test_labels_only_in_classified_cells(tmp_path):
    src = _ru_hr(tmp_path)
    out = tmp_path / "o.xlsx"
    table.apply_table(src, table.scan_table(src), Vault(), out, lang="en")
    rows = _values(out)
    assert all(LABEL.match(r[0]) and r[0].startswith("[PERSON_") for r in rows[1:])
    assert all(r[1].startswith("[ID_") for r in rows[1:])
