import docx
import pytest

from anon.text import out_path_for, read_text, write_text


def test_read_utf8_and_bom_and_cp1251(tmp_path):
    s = "Мария Ковалева\nстрока два"
    p1 = tmp_path / "a.md"
    p1.write_bytes(s.encode("utf-8"))
    p2 = tmp_path / "b.txt"
    p2.write_bytes(b"\xef\xbb\xbf" + s.encode("utf-8"))
    p3 = tmp_path / "c.txt"
    p3.write_bytes(s.encode("cp1251"))
    assert read_text(p1) == s
    assert read_text(p2) == s
    assert read_text(p3) == s


def test_write_text_utf8(tmp_path):
    p = write_text(tmp_path / "o.md", "Текст")
    assert p.read_bytes() == "Текст".encode()


def test_read_docx_paragraphs_and_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("Интервью с Марией Ковалевой")
    d.add_paragraph("")
    d.add_paragraph("Она работает в Ромашке.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "ФИО"
    t.cell(0, 1).text = "Телефон"
    t.cell(1, 0).text = "Петр Иванов"
    t.cell(1, 1).text = "+7 916 123-45-67"
    d.add_paragraph("Конец.")
    path = tmp_path / "interview.docx"
    d.save(path)
    assert read_text(path) == (
        "Интервью с Марией Ковалевой\n\n"
        "Она работает в Ромашке.\n\n"
        "ФИО\nТелефон\nПетр Иванов\n+7 916 123-45-67\n\n"
        "Конец."
    )


def test_docx_merged_cells_not_duplicated(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Объединено"
    path = tmp_path / "m.docx"
    d.save(path)
    assert read_text(path) == "Объединено"


def test_unsupported_suffix(tmp_path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF")
    with pytest.raises(ValueError):
        read_text(p)


def test_out_path_for(tmp_path):
    assert out_path_for(tmp_path / "interview.md") == tmp_path / "interview.anon.md"
    assert out_path_for(tmp_path / "interview.txt") == tmp_path / "interview.anon.txt"
    assert out_path_for(tmp_path / "interview.docx") == tmp_path / "interview.anon.md"
    assert (
        out_path_for(tmp_path / "answer.md", suffix="restored")
        == tmp_path / "answer.restored.md"
    )
