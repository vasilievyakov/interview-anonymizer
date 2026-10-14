"""Evaluate the detectors and the replacement step against the synthetic reference set.

    uv run python scripts/eval.py                   # full evaluation
    uv run python scripts/eval.py --check-fixtures  # only validate golden.json against the files
    uv run python scripts/eval.py --json out.json   # also write the results as JSON

Text fixtures: recall per kind for each single layer and for all layers together
(a golden mention counts as found when one finding covers >= 50% of its characters),
false hits (findings touching no golden mention, decoy or indirect phrase), decoys flagged.
Then `replace.apply` with a fresh vault on the all-layers findings: leaks (golden surfaces
still present in the output) and entity consistency (one golden entity, several labels).

Table fixtures: `table.scan_table` + `table.apply_table` into a temp file, then leaks,
consistency, decoys and formulas are checked by reading the cells back.

Exit code: 1 if any leak, 2 if fixtures are broken or detector modules are missing, else 0.
The detector modules are imported lazily, so --check-fixtures works without them.
"""

from __future__ import annotations

import argparse
import importlib
import json
import re
import sys
import tempfile
import traceback
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

SINGLE_LAYERS = ("rules", "natasha", "spacy")
LABEL_RE = re.compile(r"\[[^\[\]\s]{1,24}_\d+\]")
TOKEN_RE = re.compile(r"\w+|\s+|[^\w\s]")


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------


def load_golden(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def read_text_fixture(file: str) -> str:
    return (FIX / file).read_text(encoding="utf-8")


def _cell_text(wb, where: str) -> str:
    sheet, coord = where.split("!")
    v = wb[sheet][coord].value
    return "" if v is None else str(v)


def check_fixtures(golden: dict) -> list[str]:
    """Every golden, decoy and indirect record must point at its exact surface."""
    errs: list[str] = []
    valid_kinds = _valid_kinds()
    files = golden.get("files", {})
    for key in ("mentions", "decoys", "indirect"):
        if not isinstance(golden.get(key), list):
            errs.append(f"golden.json: '{key}' missing or not a list")
    if errs:
        return errs
    workbooks = {}
    texts = {}
    for file, meta in files.items():
        path = FIX / file
        if not path.exists():
            errs.append(f"{file}: listed in golden.json but missing on disk")
            continue
        if meta.get("type") == "xlsx":
            from openpyxl import load_workbook

            workbooks[file] = load_workbook(path)
        else:
            texts[file] = path.read_text(encoding="utf-8")
    for key in ("mentions", "decoys", "indirect"):
        for i, rec in enumerate(golden[key]):
            file = rec.get("file")
            label = f"{key}[{i}] {file} {rec.get('text')!r}"
            if file not in files:
                errs.append(f"{label}: file not listed in 'files'")
                continue
            if key != "indirect" and valid_kinds and rec.get("kind") not in valid_kinds:
                errs.append(
                    f"{label}: kind {rec.get('kind')!r} is not an anon.model.Kind"
                )
            if key == "mentions" and not rec.get("entity"):
                errs.append(f"{label}: no entity id")
            if file in texts:
                src = texts[file]
            elif file in workbooks:
                if not rec.get("where"):
                    errs.append(f"{label}: xlsx record without 'where'")
                    continue
                try:
                    src = _cell_text(workbooks[file], rec["where"])
                except KeyError as e:
                    errs.append(f"{label}: bad 'where' {rec['where']}: {e}")
                    continue
            else:
                continue
            s, e = rec.get("start"), rec.get("end")
            if s is not None and e is not None:
                if src[s:e] != rec["text"]:
                    errs.append(f"{label}: offsets {s}:{e} hold {src[s:e]!r}")
            elif rec["text"] not in src:
                errs.append(f"{label}: surface not found")
    for file in files:
        if not any(m["file"] == file for m in golden["mentions"]):
            errs.append(f"{file}: no golden mentions")
    return errs


def _valid_kinds() -> set[str]:
    try:
        from anon.model import Kind

        return {k.value for k in Kind}
    except Exception:
        return set()


def fixture_stats(golden: dict) -> list[str]:
    lines = []
    for file in golden["files"]:
        ms = [m for m in golden["mentions"] if m["file"] == file]
        kinds: dict[str, int] = defaultdict(int)
        for m in ms:
            kinds[m["kind"]] += 1
        nd = sum(1 for d in golden["decoys"] if d["file"] == file)
        ni = sum(1 for d in golden["indirect"] if d["file"] == file)
        ents = len({m["entity"] for m in ms})
        lines.append(
            f"{file}: {len(ms)} mentions, {ents} entities, {nd} decoys, {ni} indirect | "
            + ", ".join(f"{k} {v}" for k, v in sorted(kinds.items()))
        )
    return lines


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------


def _overlap(a0: int, a1: int, b0: int, b1: int) -> int:
    return max(0, min(a1, b1) - max(a0, b0))


def _kind_str(k) -> str:
    return getattr(k, "value", str(k))


def _word_re(surface: str) -> re.Pattern:
    return re.compile(r"(?<!\w)" + re.escape(surface) + r"(?!\w)", re.IGNORECASE)


def _tokens(s: str) -> list[tuple[str, int, int]]:
    return [(m.group(), m.start(), m.end()) for m in TOKEN_RE.finditer(s)]


def labels_for_spans(
    src: str, out: str, spans: list[tuple[int, int]]
) -> list[str | None]:
    """Which label replaced each source span. Aligns source and output token-wise."""
    ta, tb = _tokens(src), _tokens(out)
    sm = SequenceMatcher(None, [t[0] for t in ta], [t[0] for t in tb], autojunk=False)
    chunks = []  # (a_start, a_end, labels in the replacement)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal" or i2 == i1:
            continue
        a0, a1 = ta[i1][1], ta[i2 - 1][2]
        b_text = out[tb[j1][1] : tb[j2 - 1][2]] if j2 > j1 else ""
        chunks.append((a0, a1, LABEL_RE.findall(b_text)))
    result: list[str | None] = []
    for s, e in spans:
        label = None
        for a0, a1, labels in chunks:
            if not labels or not _overlap(s, e, a0, a1):
                continue
            inside = [sp for sp in spans if _overlap(sp[0], sp[1], a0, a1)]
            inside.sort()
            idx = inside.index((s, e)) if (s, e) in inside else 0
            label = labels[idx] if idx < len(labels) else labels[-1]
            break
        result.append(label)
    return result


def _load_modules(need: list[str]) -> tuple[dict, list[str]]:
    mods, missing = {}, []
    for name in need:
        try:
            mods[name] = importlib.import_module(name)
        except Exception as e:  # ImportError, or a module that fails on import
            missing.append(f"{name} ({type(e).__name__}: {e})")
    return mods, missing


# --------------------------------------------------------------------------
# Text evaluation
# --------------------------------------------------------------------------


def score_layer(text: str, findings, mentions, decoys, indirect) -> dict:
    found_by_kind: dict[str, list[int]] = defaultdict(
        lambda: [0, 0, 0]
    )  # found, total, kind ok
    missed = []
    for m in mentions:
        s, e = m["start"], m["end"]
        need = 0.5 * (e - s)
        best = None
        for f in findings:
            ov = _overlap(s, e, f.start, f.end)
            if ov >= need and (best is None or ov > best[0]):
                best = (ov, f)
        row = found_by_kind[m["kind"]]
        row[1] += 1
        if best:
            row[0] += 1
            if _kind_str(best[1].kind) == m["kind"]:
                row[2] += 1
        else:
            missed.append({"text": m["text"], "kind": m["kind"], "start": s})
    decoys_hit = []
    for d in decoys:
        s, e = d["start"], d["end"]
        if any(_overlap(s, e, f.start, f.end) >= 0.5 * (e - s) for f in findings):
            decoys_hit.append(
                {
                    "text": d["text"],
                    "kind": d["kind"],
                    "type": d.get("type"),
                    "public": d.get("public", False),
                }
            )
    false_hits = []
    protected = [(r["start"], r["end"]) for r in [*mentions, *decoys, *indirect]]
    for f in findings:
        if not any(_overlap(f.start, f.end, s, e) for s, e in protected):
            false_hits.append(
                {
                    "text": f.text,
                    "kind": _kind_str(f.kind),
                    "start": f.start,
                    "sources": list(getattr(f, "sources", [])),
                }
            )
    total_found = sum(v[0] for v in found_by_kind.values())
    total = sum(v[1] for v in found_by_kind.values())
    return {
        "recall": {
            k: {"found": v[0], "total": v[1], "kind_ok": v[2]}
            for k, v in sorted(found_by_kind.items())
        },
        "total": {"found": total_found, "total": total},
        "missed": missed,
        "false_hits": false_hits,
        "decoys_hit": decoys_hit,
        "findings": len(findings),
    }


def eval_text(file: str, golden: dict, mods: dict) -> dict:
    detect = mods["anon.detect"]
    text = read_text_fixture(file)
    mentions = [m for m in golden["mentions"] if m["file"] == file]
    decoys = [d for d in golden["decoys"] if d["file"] == file]
    indirect = [d for d in golden["indirect"] if d["file"] == file]
    res: dict = {"file": file, "layers": {}}
    all_layers = tuple(getattr(detect, "ALL_LAYERS", SINGLE_LAYERS))
    runs = [(l, (l,)) for l in SINGLE_LAYERS] + [("all", all_layers)]
    all_findings = None
    for name, layers in runs:
        try:
            findings = detect.find_all(text, layers=layers)
        except Exception as e:
            res["layers"][name] = {"error": f"{type(e).__name__}: {e}"}
            continue
        res["layers"][name] = score_layer(text, findings, mentions, decoys, indirect)
        if name == "all":
            all_findings = findings

    if all_findings is None:
        res["replace"] = {"error": "no all-layers findings"}
        return res
    try:
        from anon import lang as lang_mod

        out_lang = lang_mod.label_language(text)
        vault = mods["anon.vault"].Vault()
        out = mods["anon.replace"].apply(text, all_findings, vault, out_lang)
    except Exception as e:
        res["replace"] = {
            "error": f"{type(e).__name__}: {e}",
            "trace": traceback.format_exc(limit=3),
        }
        return res
    res["replace"] = leak_and_consistency_text(text, out, mentions, decoys)
    return res


def leak_and_consistency_text(text: str, out: str, mentions, decoys) -> dict:
    leaks = []
    for surface in sorted({m["text"] for m in mentions}, key=str.lower):
        n = len(_word_re(surface).findall(out))
        if n:
            kind = next(m["kind"] for m in mentions if m["text"] == surface)
            leaks.append({"text": surface, "kind": kind, "count": n})
    spans = [(m["start"], m["end"]) for m in mentions]
    labels = labels_for_spans(text, out, spans)
    return {
        "leaks": leaks,
        "leak_count": len(leaks),
        **_consistency(mentions, labels),
        "decoys_lost": [
            d["text"] for d in decoys if not _word_re(d["text"]).search(out)
        ],
    }


def _consistency(mentions, labels) -> dict:
    by_entity: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    unlabeled = 0
    for m, lab in zip(mentions, labels):
        if lab is None:
            unlabeled += 1
            continue
        by_entity[m["entity"]][lab].append(m["text"])
    split = []
    for ent, labs in sorted(by_entity.items()):
        if len(labs) > 1:
            split.append(
                {"entity": ent, "labels": {k: sorted(set(v)) for k, v in labs.items()}}
            )
    # one label shared by different golden entities of one kind (namesakes merged)
    owners: dict[str, set[str]] = defaultdict(set)
    for ent, labs in by_entity.items():
        for lab in labs:
            owners[lab].add(ent)
    merged = [
        {"label": lab, "entities": sorted(ents)}
        for lab, ents in sorted(owners.items())
        if len(ents) > 1
    ]
    return {
        "entities_checked": len(by_entity),
        "entities_split": split,
        "labels_merged": merged,
        "mentions_without_label": unlabeled,
    }


# --------------------------------------------------------------------------
# Table evaluation
# --------------------------------------------------------------------------


def eval_table(file: str, golden: dict, mods: dict) -> dict:
    from openpyxl import load_workbook

    table = mods["anon.table"]
    path = FIX / file
    mentions = [m for m in golden["mentions"] if m["file"] == file]
    decoys = [d for d in golden["decoys"] if d["file"] == file]
    res: dict = {"file": file}
    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / path.name
        try:
            scan = table.scan_table(path)
            vault = mods["anon.vault"].Vault()
            table.apply_table(path, scan, vault, out_path)
        except Exception as e:
            res["error"] = f"{type(e).__name__}: {e}"
            res["trace"] = traceback.format_exc(limit=3)
            return res
        wb_in = load_workbook(path)
        wb_out = load_workbook(out_path)

    leaks = []
    labels: list[str | None] = []
    for m in mentions:
        try:
            src = _cell_text(wb_in, m["where"])
            dst = _cell_text(wb_out, m["where"])
        except KeyError:
            leaks.append(
                {
                    "text": m["text"],
                    "kind": m["kind"],
                    "where": m["where"],
                    "note": "sheet missing in output",
                }
            )
            labels.append(None)
            continue
        if _word_re(m["text"]).search(dst):
            leaks.append({"text": m["text"], "kind": m["kind"], "where": m["where"]})
        labels.append(labels_for_spans(src, dst, [(m["start"], m["end"])])[0])

    decoys_lost = []
    for d in decoys:
        try:
            if d["text"] not in _cell_text(wb_out, d["where"]):
                decoys_lost.append(f"{d['where']} {d['text']}")
        except KeyError:
            decoys_lost.append(f"{d['where']} {d['text']}")

    formulas_in, formulas_lost, numbers_changed, headers_changed = 0, [], [], []
    for ws in wb_in.worksheets:
        if ws.title not in wb_out.sheetnames:
            continue
        wo = wb_out[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v, o = c.value, wo[c.coordinate].value
                if isinstance(v, str) and v.startswith("="):
                    formulas_in += 1
                    if o != v:
                        formulas_lost.append(f"{ws.title}!{c.coordinate}")
                elif isinstance(v, (int, float)) and not isinstance(v, bool) and o != v:
                    numbers_changed.append(f"{ws.title}!{c.coordinate}")
                elif c.row == 1 and v != o:
                    headers_changed.append(f"{ws.title}!{c.coordinate}")

    res.update(
        {
            "leaks": leaks,
            "leak_count": len(leaks),
            **_consistency(mentions, labels),
            "decoys_lost": decoys_lost,
            "sheets_out": wb_out.sheetnames,
            "formulas": {"in": formulas_in, "changed": formulas_lost},
            "numbers_changed": numbers_changed,
            "headers_changed": headers_changed,
        }
    )
    return res


# --------------------------------------------------------------------------
# Printing
# --------------------------------------------------------------------------


def print_text_result(r: dict, verbose: bool) -> None:
    print(f"\n== {r['file']}")
    kinds = sorted(
        {k for lr in r["layers"].values() if "recall" in lr for k in lr["recall"]}
    )
    head = ["layer"] + kinds + ["total", "false", "decoys"]
    rows = []
    for name, lr in r["layers"].items():
        if "error" in lr:
            rows.append([name, f"ERROR {lr['error']}"])
            continue
        row = [name]
        for k in kinds:
            v = lr["recall"].get(k)
            row.append(f"{v['found']}/{v['total']}" if v else "-")
        t = lr["total"]
        pct = 100 * t["found"] / t["total"] if t["total"] else 0
        row += [
            f"{t['found']}/{t['total']} {pct:.0f}%",
            str(len(lr["false_hits"])),
            str(len(lr["decoys_hit"])),
        ]
        rows.append(row)
    _table(head, rows)

    lr = r["layers"].get("all", {})
    if "recall" in lr:
        if lr["missed"]:
            print(f"missed by all layers ({len(lr['missed'])}):")
            for m in lr["missed"][: None if verbose else 40]:
                print(f"   {m['kind']:<9} {m['text']}")
        if lr["false_hits"]:
            print(f"false hits, all layers ({len(lr['false_hits'])}):")
            for f in lr["false_hits"]:
                print(f"   {f['kind']:<9} {f['text']!r}  [{','.join(f['sources'])}]")
        if lr["decoys_hit"]:
            print(f"decoys flagged, all layers ({len(lr['decoys_hit'])}):")
            for d in lr["decoys_hit"]:
                pub = " (public, reviewer may keep)" if d["public"] else ""
                print(f"   {d['type']:<9} {d['text']!r}{pub}")
    _print_replace(r.get("replace", {}))


def _print_replace(rep: dict) -> None:
    if "error" in rep:
        print(f"replace: ERROR {rep['error']}")
        return
    print(f"leaks after replace: {rep['leak_count']}")
    for l in rep["leaks"]:
        where = f" at {l['where']}" if "where" in l else ""
        cnt = f" x{l['count']}" if "count" in l else ""
        print(f"   LEAK {l['kind']:<9} {l['text']!r}{cnt}{where}")
    print(
        f"consistency: {rep['entities_checked']} entities labelled, "
        f"{len(rep['entities_split'])} split across labels, "
        f"{len(rep['labels_merged'])} labels shared by different entities, "
        f"{rep['mentions_without_label']} mentions without a label"
    )
    for s in rep["entities_split"]:
        parts = "; ".join(f"{lab}: {', '.join(v)}" for lab, v in s["labels"].items())
        print(f"   SPLIT {s['entity']}: {parts}")
    for m in rep["labels_merged"]:
        print(f"   MERGED {m['label']}: {', '.join(m['entities'])}")
    if rep.get("decoys_lost"):
        print(
            f"decoys replaced in output ({len(rep['decoys_lost'])}): {', '.join(map(repr, rep['decoys_lost'][:20]))}"
        )


def print_table_result(r: dict) -> None:
    print(f"\n== {r['file']}")
    if "error" in r:
        print(f"ERROR {r['error']}")
        return
    _print_replace(r)
    f = r["formulas"]
    print(
        f"formulas: {f['in']} in source, {len(f['changed'])} changed; numbers changed: {len(r['numbers_changed'])}; headers changed: {len(r['headers_changed'])}"
    )


def _table(head: list[str], rows: list[list[str]]) -> None:
    widths = [
        max(len(str(x[i])) for x in [head, *rows] if i < len(x))
        for i in range(len(head))
    ]
    fmt = lambda row: "  ".join(str(c).ljust(widths[i]) for i, c in enumerate(row))  # noqa: E731
    print(fmt(head))
    for row in rows:
        print(fmt(row))


# --------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--golden", type=Path, default=FIX / "golden.json")
    ap.add_argument(
        "--check-fixtures",
        action="store_true",
        help="only validate golden.json against the fixture files",
    )
    ap.add_argument("--json", type=Path, help="write results as JSON to this path")
    ap.add_argument(
        "--only", help="evaluate only fixture files containing this substring"
    )
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    if not args.golden.exists():
        print(
            f"golden file not found: {args.golden}. Run: uv run python scripts/make_fixtures.py"
        )
        return 2
    golden = load_golden(args.golden)
    errs = check_fixtures(golden)
    if errs:
        print(f"fixture check FAILED, {len(errs)} problems:")
        for e in errs:
            print("  ", e)
        return 2
    print("fixture check OK")
    for line in fixture_stats(golden):
        print("  ", line)
    if args.check_fixtures:
        if args.json:
            args.json.write_text(
                json.dumps({"fixtures_ok": True}, indent=1), encoding="utf-8"
            )
        return 0

    files = [f for f in golden["files"] if not args.only or args.only in f]
    text_files = [f for f in files if golden["files"][f]["type"] == "text"]
    table_files = [f for f in files if golden["files"][f]["type"] == "xlsx"]

    results: dict = {"text": [], "tables": [], "missing_modules": []}
    text_mods, miss_text = _load_modules(["anon.detect", "anon.replace", "anon.vault"])
    table_mods, miss_table = _load_modules(["anon.table", "anon.vault"])
    if miss_text and text_files:
        print("\ntext evaluation skipped, detector modules not available yet:")
        for m in miss_text:
            print("   ", m)
        results["missing_modules"] += miss_text
    else:
        for f in text_files:
            r = eval_text(f, golden, text_mods)
            results["text"].append(r)
            print_text_result(r, args.verbose)
    if miss_table and table_files:
        print("\ntable evaluation skipped, modules not available yet:")
        for m in miss_table:
            print("   ", m)
        results["missing_modules"] += [
            m for m in miss_table if m not in results["missing_modules"]
        ]
    else:
        for f in table_files:
            r = eval_table(f, golden, table_mods)
            results["tables"].append(r)
            print_table_result(r)

    leaks = sum(r.get("replace", {}).get("leak_count", 0) for r in results["text"])
    leaks += sum(r.get("leak_count", 0) for r in results["tables"])
    errors = [r["file"] for r in results["text"] if "error" in r.get("replace", {})]
    errors += [r["file"] for r in results["tables"] if "error" in r]
    results["leaks_total"] = leaks
    results["errors"] = errors
    print(
        f"\nTOTAL leaks: {leaks}"
        + (f"; runs with errors: {', '.join(errors)}" if errors else "")
    )
    if args.json:
        args.json.write_text(
            json.dumps(results, ensure_ascii=False, indent=1, default=str),
            encoding="utf-8",
        )
        print(f"wrote {args.json}")
    if leaks:
        return 1
    if results["missing_modules"] or errors:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
