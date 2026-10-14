"""Generate the synthetic reference set: tests/fixtures/{ru,en}/... and golden.json.

All data is invented. Identifiers (SNILS, INN, SSN, NINO, card) are produced here
with valid checksums or formats from a fixed seed, so the output is deterministic.

Text templates use inline markup that is stripped when the file is written:

    [[P:entity|surface]]         golden mention (P, O, L, PH, E, S, INN, PS, C, SSN, NI, ID)
    [[d.KIND.type|surface]]      decoy that must stay; append ".pub" for public platforms
    [[i|surface]]                indirect identifier (not expected to be auto-detected)

Markup may nest (an indirect phrase can contain a golden LOCATION).
`<<NAME>>` placeholders are filled with generated identifiers before parsing.

Run:  uv run python scripts/make_fixtures.py
"""

from __future__ import annotations

import json
import random
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"

KINDS = {
    "P": "PERSON",
    "O": "ORG",
    "L": "LOCATION",
    "PH": "PHONE",
    "E": "EMAIL",
    "S": "SNILS",
    "INN": "INN",
    "PS": "PASSPORT",
    "C": "CARD",
    "SSN": "SSN",
    "NI": "NINO",
    "ID": "ID",
}
VALID_KINDS = set(KINDS.values())


class FixtureError(Exception):
    pass


# --------------------------------------------------------------------------
# Identifier generators (invented values, valid checksums)
# --------------------------------------------------------------------------


def snils(rng: random.Random) -> str:
    while True:
        n = rng.randint(2_000_000, 999_999_999)
        d = [int(c) for c in f"{n:09d}"]
        # SNILS forbids three identical digits in a row
        if any(d[i] == d[i + 1] == d[i + 2] for i in range(7)):
            continue
        s = sum(d[i] * (9 - i) for i in range(9))
        if s < 100:
            c = s
        elif s in (100, 101):
            c = 0
        else:
            c = s % 101
            if c == 100:
                c = 0
        t = f"{n:09d}"
        return f"{t[0:3]}-{t[3:6]}-{t[6:9]} {c:02d}"


def snils_valid(v: str) -> bool:
    digits = re.sub(r"\D", "", v)
    if len(digits) != 11:
        return False
    d = [int(c) for c in digits[:9]]
    s = sum(d[i] * (9 - i) for i in range(9))
    c = s if s < 100 else (0 if s in (100, 101) else s % 101 % 100)
    return c == int(digits[9:])


def _inn_ctrl(d: list[int], w: list[int]) -> int:
    return sum(a * b for a, b in zip(d, w)) % 11 % 10


def inn10(rng: random.Random) -> str:
    d = [rng.randint(1, 9)] + [rng.randint(0, 9) for _ in range(8)]
    d.append(_inn_ctrl(d, [2, 4, 10, 3, 5, 9, 4, 6, 8]))
    return "".join(map(str, d))


def inn12(rng: random.Random) -> str:
    d = [rng.randint(1, 9)] + [rng.randint(0, 9) for _ in range(9)]
    d.append(_inn_ctrl(d, [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]))
    d.append(_inn_ctrl(d, [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]))
    return "".join(map(str, d))


def luhn_card(rng: random.Random, prefix: str = "400012") -> str:
    body = prefix + "".join(str(rng.randint(0, 9)) for _ in range(15 - len(prefix)))
    total = 0
    for i, ch in enumerate(reversed(body)):
        x = int(ch)
        if i % 2 == 0:  # these positions are doubled once the check digit is appended
            x *= 2
            if x > 9:
                x -= 9
        total += x
    num = body + str((10 - total % 10) % 10)
    return " ".join(num[i : i + 4] for i in range(0, 16, 4))


def luhn_valid(v: str) -> bool:
    digits = [int(c) for c in re.sub(r"\D", "", v)]
    total = 0
    for i, x in enumerate(reversed(digits)):
        if i % 2 == 1:
            x *= 2
            if x > 9:
                x -= 9
        total += x
    return total % 10 == 0


def ssn(rng: random.Random) -> str:
    while True:
        area = rng.randint(1, 899)
        if area == 666:
            continue
        group = rng.randint(1, 99)
        serial = rng.randint(1, 9999)
        return f"{area:03d}-{group:02d}-{serial:04d}"


def nino(rng: random.Random) -> str:
    first = "ABCEGHJKLMNOPRSTWXYZ"
    second = "ABCEGHJKLMNPRSTWXYZ"
    bad = {"BG", "GB", "KN", "NK", "NT", "TN", "ZZ"}
    while True:
        p = rng.choice(first) + rng.choice(second)
        if p in bad:
            continue
        digits = f"{rng.randint(0, 999999):06d}"
        return f"{p} {digits[0:2]} {digits[2:4]} {digits[4:6]} {rng.choice('ABCD')}"


# --------------------------------------------------------------------------
# Markup parser
# --------------------------------------------------------------------------

HEADER_RE = re.compile(r"^[A-Za-z0-9_.:]+$")


def render(
    template: str, subs: dict[str, str] | None = None
) -> tuple[str, list[tuple[str, int, int]]]:
    """Strip markup. Returns plain text and (header, start, end) marks."""
    for k, v in (subs or {}).items():
        template = template.replace(f"<<{k}>>", v)
    if "<<" in template:
        raise FixtureError(
            f"unfilled placeholder near: {template[template.index('<<') :][:40]!r}"
        )
    out: list[str] = []
    stack: list[tuple[str, int]] = []
    marks: list[tuple[str, int, int]] = []
    i = 0
    n = len(template)
    while i < n:
        if template.startswith("[[", i):
            j = template.find("|", i)
            header = template[i + 2 : j]
            if j < 0 or not HEADER_RE.match(header):
                raise FixtureError(
                    f"bad markup header at {i}: {template[i : i + 40]!r}"
                )
            stack.append((header, len(out)))
            i = j + 1
            continue
        if template.startswith("]]", i) and stack:
            header, start = stack.pop()
            marks.append((header, start, len(out)))
            i += 2
            continue
        out.append(template[i])
        i += 1
    if stack:
        raise FixtureError(f"unclosed markup: {stack}")
    return "".join(out), marks


def classify(file: str, text: str, marks, base: dict | None = None):
    """Turn marks into golden / decoy / indirect records."""
    golden, decoys, indirect = [], [], []
    base = base or {}
    for header, s, e in marks:
        surface = text[s:e]
        if header == "i":
            indirect.append(
                {"file": file, "text": surface, "start": s, "end": e, **base}
            )
        elif header.startswith("d."):
            parts = header.split(".")
            kind, dtype = parts[1], parts[2]
            rec = {
                "file": file,
                "text": surface,
                "kind": kind,
                "type": dtype,
                "start": s,
                "end": e,
                **base,
            }
            if "pub" in parts[3:]:
                rec["public"] = True
            if kind not in VALID_KINDS:
                raise FixtureError(f"decoy kind {kind} invalid in {header}")
            decoys.append(rec)
        else:
            code, entity = header.split(":", 1)
            if code not in KINDS:
                raise FixtureError(f"unknown kind code {code} in {header}")
            golden.append(
                {
                    "file": file,
                    "text": surface,
                    "kind": KINDS[code],
                    "entity": entity,
                    "start": s,
                    "end": e,
                    **base,
                }
            )
    return golden, decoys, indirect


# --------------------------------------------------------------------------
# Russian interview
# --------------------------------------------------------------------------

RU_INTERVIEW = r"""# Кастдев: селлер, импорт товаров для дома

Дата записи: [[d.ID.date|12 марта 2026]]. Формат: созвон, расшифровка с таймкодами, речь не редактировалась.
Респондент: [[P:gusev|Гусев Денис Андреевич]], [[O:ip_gusev|ИП Гусев Д.А.]]

---

\[00:00:14\] **[[P:gusev|Денис]], привет. Спасибо, что нашел время. Расскажи коротко, чем ты занимаешься?**

ну привет. короче, я вожу товары из [[d.LOCATION.country.pub|Китая]] и продаю на маркетплейсах, в основном на [[d.ORG.platform.pub|Wildberries]] и немножко на [[d.ORG.platform.pub|Ozon]]. оформлен как [[O:ip_gusev|ИП Гусев]], но по факту нас уже человек семь, так что это скорее такая маленькая контора, чем ИП. сам я из [[L:ivanovo|Иванова]], но последние три года живу в [[L:moscow|Москве]], склад у нас под [[L:podolsk|Подольском]]. торгуем термокружками, органайзерами для кухни, силиконовыми формами, ну и всякой мелочью для дома, что заходит. сезонное тоже берем, к новому году гирлянды, летом всякие штуки для дачи.

\[00:01:02\] **А как ты вообще пришел в эту тему?**

да случайно, если честно. в [[d.ID.year|2021]] году мой одноклассник, [[P:belov|Санек]], ну [[P:belov|Саня Белов]], он уже продавал на [[d.ORG.platform.pub|вб]] чехлы для телефонов, и как-то сказал: давай, говорит, попробуй, там деньги легкие. легкие, ага. я вложил первые [[d.ID.amount|200 тысяч]], половину слил на рекламу, вторую на товар, который никто не купил. это были такие подставки под ноутбук, их на тот момент продавали все кому не лень. потом уже разобрался, начал смотреть аналитику, считать. [[P:belov|Санек]], кстати, до сих пор помогает, мы с ним раз в месяц созваниваемся, обсуждаем, кто что везет, у кого какие проблемы на приемке. он такой, знаешь, человек-справочник, всех знает.

\[00:02:33\] **Кто у тебя сейчас в команде, кто за что отвечает?**

ну смотри. бухгалтерию ведет [[P:kovaleva|Мария Ковалева]], она формально на аутсорсе, но фактически только с нами и работает. с [[P:kovaleva|Марией Ковалевой]] мы вообще с первого дня, она мне и ИП открывать помогала, и налоговую объясняла, я ж вообще ничего не понимал тогда. склад держит [[P:petrov|Сергей Петров]], у [[P:petrov|Петрова]] там двое ребят на сборке и упаковке. закупками в [[d.LOCATION.country.pub|Китае]] занимается [[P:morozov_o|Олег Викторович Морозов]], он живет в [[L:guangzhou|Гуанчжоу]] уже лет десять, через него идет вся связь с фабриками. ну и я, на мне карточки, реклама, деньги, переговоры, все остальное. то есть по сути я и директор, и маркетолог, и уборщица иногда.

\[00:03:40\] **Расскажи подробнее про закупки. Как устроена работа с [[P:morozov_o|Олегом Викторовичем]]?**

с [[P:morozov_o|Олегом Викторовичем]] все строится на доверии, честно. я ему пишу, что хочу, кидаю ссылки, фотки, он едет на фабрику, смотрит образцы, торгуется. обычно минимальная партия где-то на [[d.ID.amount|5000 юаней]], иногда больше, если товар крупный. он берет свои [[d.ID.percent|5%]] от закупки. [[P:morozov_o|Викторович]] человек старой школы, он не любит переписку, ему проще позвонить и все голосом обсудить. его китайский номер я даже не помню, звоню на российский, [[PH:morozov_o_phone|+7 (916) 482-35-17]], он его держит в роуминге. ну и бывает, что он пропадает на неделю, и ты сидишь как дурак, не знаешь, отгрузили или нет. а у тебя уже реклама запущена, карточки висят, люди заказывают.

\[00:05:12\] **А что самое больное во всей этой цепочке?**

логистика, конечно. карго. мы раньше работали с [[O:vostok_cargo|АО «Восток-Карго Логистик»]], потом ушли, потому что они у нас потеряли партию на [[d.ID.amount|60000]] рублей... нет, вру, на [[d.ID.amount|600 тысяч]]. там была целая история, я с их менеджером [[P:morozov_i|Игорем Морозовым]] полгода переписывался. кстати, не путай, это другой [[P:morozov_i|Морозов]], не наш [[P:morozov_o|Олег Викторович]], просто однофамилец, они друг друга даже не знают, я проверял. так вот, [[P:morozov_i|Игорь]] все время говорил, что груз на таможне, потом что на складе в [[L:zabaikalsk|Забайкальске]], потом вообще перестал отвечать. в итоге вернули [[d.ID.percent|40 процентов]] стоимости, и то через претензию и через юриста.

\[00:07:05\] **И к кому вы ушли после этого?**

сейчас возим через [[O:baikal|ООО «Байкал-Транзит»]], там директор [[P:safonova|Ирина Сафонова]], очень толковая женщина. с [[P:safonova|Ириной]] можно договориться, она всегда на связи, даже в выходные отвечает. я тебе могу ее контакт дать, если хочешь с ней тоже поговорить, вот почта [[E:safonova_email|i.safonova@example.com]], и мобильный [[PH:safonova_phone|8-903-517-26-44]]. только скажи, что от меня, от [[P:gusev|Дениса]], она тогда нормально ответит. у них, кстати, договор был на юрлицо, ИНН [[INN:baikal_inn|<<INN_BAIKAL>>]], я его наизусть уже помню, столько раз в платежках вбивал.

\[00:08:21\] **Спасибо. А как вы считаете экономику одной партии?**

ну тут все в таблице у [[P:kovaleva|Маши]]. она считает себестоимость с доставкой, пошлиной, сертификацией, комиссией маркетплейса. кружка, допустим, обходится в [[d.ID.amount|180 рублей]] с доставкой, продаем за [[d.ID.amount|990]], но там комиссия [[d.ID.percent|25%]], логистика маркетплейса, возвраты, реклама. в чистую выходит рублей сто пятьдесят, двести со штуки, если повезет. в прошлом месяце оборот был где-то [[d.ID.amount|3,2 миллиона]], а в кармане осталось тысяч триста. и это еще хороший месяц был, в январе вообще в ноль ушли. я, честно, иногда сам не понимаю, куда деньги деваются, пока [[P:kovaleva|Маша]] не сядет и не разложит все по полочкам.

\[00:10:02\] **Ты упомянул сертификацию. Как с ней обстоят дела?**

ой, это отдельная боль. на каждую партию нужна декларация соответствия, мы делаем через посредника, [[O:standart|ООО "СтандартЭксперт"]]. оформляли сначала на меня как на ИП, там нужны были мои данные, паспорт, СНИЛС, ИНН, все это. я, помню, диктовал по телефону девушке: паспорт [[PS:gusev_pass|45 19 702318]], СНИЛС [[S:gusev_snils|<<SNILS_GUSEV>>]], ИНН [[INN:gusev_inn|<<INN_GUSEV>>]]. потом думаю, блин, а кому я это все диктую вообще, я ж даже не знаю, кто на том конце. ну ладно, вроде нормальные ребята оказались, декларацию сделали за неделю, взяли [[d.ID.amount|18 тысяч]].

\[00:11:40\] **А с маркетплейсами как отношения? Что раздражает больше всего?**

штрафы. просто штрафы за все подряд. на прошлой неделе, [[d.ID.date|5 марта]], нам выписали штраф [[d.ID.amount|15 000]] за то, что коробка была на два сантиметра выше нормы. два сантиметра, [[d.PERSON.meme|Карл]]. и ты пишешь в поддержку, а там бот, который отвечает шаблоном. потом еще приемка: привозишь на склад, а у них нет слота, и машина стоит, водитель злится, ты платишь простой. [[d.ORG.platform.pub|Ozon]] в этом плане адекватнее, но там продаж раз в пять меньше, так что выбора особо нет. ну и правила меняются каждый месяц, вот это вообще убивает, ты только подстроился, а уже все по-другому.

\[00:13:15\] **Как ты сейчас ищешь новые товары?**

смотрю аналитику, сервисы всякие, где видно, что растет. плюс [[P:morozov_o|Олег Викторович]] присылает, что видит на выставках в [[L:guangzhou|Гуанчжоу]], там же [[d.ORG.event.pub|Кантонская ярмарка]] два раза в год. в апреле он туда пойдет, я ему уже список составил, штук тридцать позиций. ну и [[P:belov|Санек]] иногда подкидывает идеи, он больше по электронике, но глаз у него наметанный. вообще самое сложное не найти товар, а угадать, когда его брать. приходит партия через два месяца, а ниша уже забита, все такие же умные. было так с бутылками для воды, мы опоздали буквально на три недели.

\[00:15:00\] **У вас есть офис или все работают удаленно?**

есть маленький офис в [[L:moscow|Москве]], на самом деле это просто комната в коворкинге, там я и сижу, иногда [[P:kovaleva|Маша]] заезжает. и есть пункт в [[L:tver|Твери]], небольшой, там у нас сидит девочка на первичке, ну, документы, накладные. она, по сути, [[i|единственный бухгалтер в нашем филиале в [[L:tver|Твери]]]], [[P:kovaleva|Маша]] ее курирует удаленно. почему [[L:tver|Тверь]]? да просто так сложилось, там аренда дешевая и человек хороший нашелся. склад, как я говорил, под [[L:podolsk|Подольском]], там [[P:petrov|Петров]] со своими ребятами.

\[00:16:30\] **Как вы общаетесь внутри команды? Какие инструменты используете?**

телеграм, конечно. общий чат, туда все кидают, и фотки, и счета, и мемы. [[P:petrov|Петров]] скидывает фотки приемки, если что-то битое пришло. бывает, что [[P:petrov|Сергей]] пишет в два ночи, мол, фура пришла, а там половина коробок мятые. [[P:kovaleva|Маша]] больше любит почту, у нее рабочий ящик [[E:kovaleva_email|kovaleva.m@example.org]], она туда все акты и счета собирает, говорит, что в телеграме все теряется. ну и гугл-таблицы, куда без них. мой номер ты знаешь, но если вдруг потеряешь, [[PH:gusev_phone|89265550147]], это основной, второй я почти не беру.

\[00:18:05\] **Были конфликты с подрядчиками, которые сильно ударили по бизнесу?**

ну вот эта история с карго главная. когда потерялась партия, я реально думал, что все, закрываемся. кредит висит, товара нет, а в карточках уже отзывы копятся. нам тогда помог юрист, [[P:levin|Аркадий Левин]] из [[O:pravovoy|ООО «Правовой Щит»]], он писал претензию, потом исковое готовил, но до суда не дошло. [[P:levin|Левину]] мы заплатили [[d.ID.amount|45 тысяч]], и это были лучшие потраченные деньги за тот год. он, кстати, сказал, что у [[O:vostok_cargo|«Восток-Карго»]] таких историй вагон, просто люди не судятся, боятся. а еще был случай с фабрикой, которая прислала кружки без крышек. [[P:morozov_o|Олег Викторович]] тогда сам туда поехал и полдня с ними ругался.

\[00:19:44\] **Кто у вас принимает финансовые решения?**

я. ну, формально я, но без [[P:kovaleva|Ковалевой]] ни одна платежка не уходит. в документах она у меня записана как [[P:kovaleva|Ковалева М.А.]], ответственная за учет, у нее доступ к банку. [[P:kovaleva|Мария Андреевна]], я ее так зову, когда она меня ругает за лишние траты. для закупок у меня отдельная карта, [[C:gusev_card|<<CARD_RU>>]], на нее [[P:morozov_o|Олег Викторович]] выставляет счета, а [[P:kovaleva|Маша]] потом сверяет. крупные вещи, типа новой партии на миллион, мы обсуждаем втроем, я, [[P:kovaleva|Маша]] и [[P:petrov|Сергей]], потому что ему потом это все на складе размещать.

\[00:21:10\] **Сколько времени в неделю уходит на рутину, которую хочется отдать?**

ой, много. часов двадцать, наверное. отчеты, сверки, ответы на отзывы, карточки переписывать. отзывы вообще отдельная история, их десятки в день, и на каждый надо как-то ответить, иначе рейтинг падает. я пробовал отдавать это фрилансеру, но он отвечал так, что потом еще хуже было. сейчас частично отвечаю сам, частично нейросеть помогает. отчеты [[P:kovaleva|Маша]] делает, но ей тоже не разорваться, у нее кроме нас еще пара клиентов по мелочи.

\[00:22:30\] **Ты пробовал отдавать отчеты или таблицы в нейросеть?**

пробовал, да. и вот тут у меня, честно, паранойя. я как-то закинул в чат-бот таблицу с зарплатами, чтобы он мне посчитал фонд на квартал. а там все фамилии, СНИЛСы, оклады. он посчитал, все красиво, а потом я сижу и думаю: а это все теперь где? на каких серверах? [[P:kovaleva|Маша]], когда узнала, меня чуть не убила. говорит, ты понимаешь, что это персональные данные, за это штрафы. у [[P:belov|Александра Белова]], у [[P:belov|Санька]] то есть, вообще была история, он загрузил переписку с поставщиком, а там паспорт поставщика лежал во вложении. никто, конечно, ничего не узнал, но осадочек остался.

\[00:24:15\] **Что бы ты хотел, чтобы такой инструмент умел?**

ну, чтобы я закинул таблицу или текст, а он сам все фамилии, телефоны, номера убрал, и отдал мне чистое. а потом, когда нейросеть ответила, обратно подставил, чтобы я не сидел и не расшифровывал, кто такой [[d.PERSON.label|ИМЯ_3]]. и чтобы все это у меня на ноутбуке работало, без всяких облаков, иначе какой смысл. и чтобы он не ломал цифры. вот если он мне вместо [[d.ID.amount|200000]] рублей поставит какую-то метку, то мне такой инструмент не нужен, мне цифры как раз и нужны. и даты тоже, мне важно, что партия пришла [[d.ID.date|14 февраля]], а не просто когда-то.

\[00:25:40\] **Как вы нанимаете людей? Где их ищете?**

в основном через знакомых, если честно. объявления пробовали, приходят люди, которые через месяц уходят. [[P:petrov|Сергея]] мне, например, посоветовал [[P:belov|Санек]], они раньше вместе на одном складе работали. а последнюю девочку на маркетплейсы, [[P:lapteva|Веру Лаптеву]], нашли через знакомых в [[L:kazan|Казани]], она работает удаленно, карточки делает, фотки обрабатывает. [[P:lapteva|Вера]] вообще находка, у нее глаз дизайнерский, карточки после нее продают процентов на двадцать лучше. оформляем всех через [[P:kovaleva|Машу]], она знает, какие бумаги нужны, я туда даже не лезу. ну и испытательный срок месяц, если человек за месяц не втянулся, то уже и не втянется, проверено.

\[00:27:10\] **Где у вас хранятся документы сотрудников?**

ох. ну вот тут, наверное, стыдно признаваться. часть в облаке, в папке «Кадры», часть у [[P:kovaleva|Маши]] на компьютере, часть просто в телеграме. когда [[P:lapteva|Вера]] устраивалась, она скинула скан паспорта прямо в общий чат, и он там до сих пор висит, наверное, никто не удалил. трудовые договоры подписывает [[P:kovaleva|Мария Андреевна]], у нее и оригиналы. я понимаю, что это не очень правильно, но когда вас семь человек, как-то не думаешь про регламенты. вот если бы кто-то пришел и сказал, вот так надо, вот кнопка, я бы сделал. а самому разбираться некогда, у меня каждый день то приемка, то возвраты.

\[00:28:35\] **Бывали у вас проверки?**

было один раз, в [[d.ID.year|2024]] году. пришло требование из [[d.ORG.gov.pub|ФНС]], хотели пояснения по возвратам, там что-то не сходилось по выручке. [[P:kovaleva|Маша]] два дня сидела, собирала документы, в итоге все сошлось, но штраф [[d.ID.amount|30 тысяч]] все равно заплатили, за то, что поздно сдали какой-то отчет. после этого я ей поднял зарплату, честно. а еще была история с маркировкой, когда товар без кодов приехал, но это уже маркетплейс штрафовал, а не государство. в общем, нервов много, денег тоже, а пользы от этого никакой.

\[00:30:05\] **Как ты видишь бизнес через год?**

хочу открыть второй склад, в [[L:kazan|Казани]], поближе к [[P:lapteva|Вере]] и вообще к регионам, там логистика дешевле получается. и жена, [[P:guseva|Наталья]], хочет свое направление, детские товары, говорит, что я в этом ничего не понимаю, и она права. [[P:guseva|Наташа]] уже ищет поставщиков, тоже через [[P:morozov_o|Олега Викторовича]]. если получится, будет второе ИП, на нее. [[P:belov|Санек]] говорит, что я сошел с ума, но он так про все мои идеи говорит, а потом сам повторяет. ну и хочется наконец навести порядок в документах, чтобы не стыдно было.

\[00:31:20\] **Кому еще из твоих знакомых селлеров это может быть интересно?**

[[P:belov|Саньку]], однозначно, он сам об этом говорил. еще есть [[P:kuznetsova|Оля Кузнецова]] из [[L:kazan|Казани]], она возит текстиль, у нее [[O:textile|ООО «Текстиль-Дом»]], человек двадцать сотрудников, там вообще вся кадровая история в таблицах. ее телефон [[PH:kuzn_phone|+7 917 264 83 05]], можешь сказать, что от меня. и я бы еще поговорил с [[i|девушкой из закупок [[O:baikal|«Байкал-Транзита»]], которая в прошлом году выиграла у них внутренний конкурс]], она очень системная, ей такое понравится. а, и [[i|наш кладовщик, который в прошлом году выиграл чемпионат области по армрестлингу]], он ведет учет смен, тоже в таблицах, и тоже, наверное, все туда пихает.

\[00:33:05\] **Последний вопрос. Сколько ты был бы готов за это платить?**

ну смотря как работает. если реально экономит время и я не боюсь за данные, то тысячи две-три в месяц спокойно. если будет на [[d.ORG.platform.pub|Ozon]] или на [[d.ORG.platform.pub|Wildberries]] интеграция, чтобы отчеты сразу чистые выгружать, то и [[d.ID.amount|5000]] нормально. для сравнения, я на аналитику трачу где-то [[d.ID.amount|12 тысяч]] в месяц, и это я еще экономлю.

\[00:33:50\] **Спасибо, [[P:gusev|Денис]]. Можно будет вернуться к тебе через месяц, показать прототип?**

да, конечно. пиши в телеграм или звони на [[PH:gusev_phone|+7 926 555-01-47]], это тот же номер, просто я его по-разному диктую. только не в понедельник, понедельник у нас приемка, я весь день на складе с [[P:petrov|Петровым]]. и [[P:safonova|Ирине Сафоновой]] напиши, правда, она тебе много интересного расскажет про логистику.
"""

# --------------------------------------------------------------------------
# English interview
# --------------------------------------------------------------------------

EN_INTERVIEW = r"""# Customer interview: home goods importer

Recorded: [[d.ID.date|18 March 2026]]. Video call, transcript with timestamps, lightly cleaned.
Interviewee: [[P:hadley|Tom Hadley]], director of [[O:hadley_ltd|Hadley & Finch Home Ltd]]

---

\[00:00:09\] **[[P:hadley|Tom]], thanks for making the time. Could you start by telling me what the business does?**

yeah, sure. so basically we import kitchen and home stuff from [[d.LOCATION.country.pub|China]] and sell it online, mostly on [[d.ORG.platform.pub|Amazon]], and we've got our own little store on [[d.ORG.platform.pub|Shopify]] as well. the company's [[O:hadley_ltd|Hadley & Finch Home Ltd]], we're based in [[L:leeds|Leeds]], warehouse is just outside [[L:bradford|Bradford]]. and about two years ago we opened a US arm, [[O:hadley_llc|Hadley Home Goods LLC]], registered in [[L:austin|Austin]], because the US market is just, you know, ten times bigger. bamboo chopping boards, spice racks, those collapsible silicone bowls, that kind of thing. nothing glamorous but it sells.

\[00:01:20\] **How did you get into it in the first place?**

honestly, by accident. I was working in logistics, hated it, and a mate of mine, [[P:ashdown|Robbo]], well, [[P:ashdown|Rob Ashdown]], he was already flipping phone cases on [[d.ORG.platform.pub|eBay]]. he said, mate, it's easy money. easy, right. I put in my first [[d.ID.amount|£12,000]] in [[d.ID.date|March 2021]], blew half of it on ads and the other half on laptop stands that nobody wanted. took me a good year to figure out what I was doing. [[P:ashdown|Robbo]] runs our warehouse now, actually, so it all worked out. he still reminds me about the laptop stands every Christmas party.

\[00:02:41\] **Who's on the team now, and who does what?**

so, the books are done by [[P:whitcombe|Sarah Whitcombe]], she's technically freelance but she's basically ours. honestly, without [[P:whitcombe|Sarah's]] spreadsheets we'd have gone under in the first year. [[P:ashdown|Rob]] runs the warehouse with two lads on packing. sourcing in [[d.LOCATION.country.pub|China]] goes through [[P:lin|Wei Lin]], he's an agent based in [[L:shenzhen|Shenzhen]], been there about twelve years, deals with all the factories for us. on the US side we've got a bookkeeper too, [[P:morrison_k|Karen Morrison]], she does the LLC's taxes. and me, I do listings, ads, cash, arguing with suppliers, all the rest of it. director, marketer, occasionally the cleaner.

\[00:04:02\] **Tell me more about sourcing. How does working with [[P:lin|Wei]] actually go?**

it's all trust, really. I send [[P:lin|Wei]] what I want, links, photos, rough target price, he goes out to the factory, checks samples, haggles. usually the minimum order is about [[d.ID.amount|$4,500]], sometimes more for bigger items. he takes [[d.ID.percent|5 per cent]] of the order value. [[P:lin|Wei's]] old school, doesn't like email much, he'd rather just ring you. I've got his UK number saved because he's over here a few times a year, [[PH:lin_phone|07700 900417]]. the frustrating bit is when he goes quiet for a week and you've no idea whether the goods have shipped, and meanwhile your ads are running and orders are coming in.

\[00:05:33\] **What's the most painful part of the whole chain?**

freight, without a doubt. we used to use [[O:pellow|Pellow Freight Ltd]], then we left because they lost a container on us, roughly [[d.ID.amount|40k]] worth of stock. I spent six months emailing their account manager, [[P:morrison_d|Dave Morrison]]. and before you ask, no relation to [[P:morrison_k|Karen]], it's just a common name, they've never met. different [[P:morrison_d|Morrison]] entirely. anyway, [[P:morrison_d|Dave]] kept telling me it was stuck in customs, then that it was sitting in a depot in [[L:felixstowe|Felixstowe]], then he just stopped replying. in the end they refunded [[d.ID.percent|30%]] and only after a formal complaint and a solicitor's letter.

\[00:07:15\] **Who did you move to after that?**

we're with [[O:northgate|Northgate Shipping Ltd]] now, the operations lead is [[P:okafor|Grace Okafor]], really sharp. you can actually talk to [[P:okafor|Grace]], she answers on weekends even. I can give you her details if you want to interview her as well, email's [[E:okafor_email|grace.okafor@example.com]] and her mobile is [[PH:okafor_phone|+44 7700 900982]]. tell her [[P:hadley|Tom]] sent you and she'll pick up. their office number is [[PH:northgate_phone|020 7946 0381]] but nobody ever answers that.

\[00:08:30\] **How do you work out the economics of a single order?**

that's all in [[P:whitcombe|Sarah's]] spreadsheet. landed cost with freight, duty, certification, marketplace fees, all of it. so a chopping board might cost us [[d.ID.amount|£3.20]] landed, we sell it at [[d.ID.amount|£19.99]], but then there's the [[d.ID.percent|15%]] referral fee, fulfilment fees, returns, ads. we clear maybe three or four quid a unit on a good day. last month turnover was about [[d.ID.amount|£210,000]] and what actually stayed in the bank was more like twenty grand. January we basically broke even. I honestly don't know where the money goes until [[P:whitcombe|Sarah]] sits me down and walks me through it.

\[00:10:05\] **You mentioned the US company. How was setting that up?**

painful. lots of forms. for the LLC they wanted my personal details for the bank and the tax side, so I was reading my SSN out over the phone to some bloke at the formation agency, [[SSN:hadley_ssn|<<SSN_HADLEY>>]], and my UK National Insurance number as well, [[NI:hadley_nino|<<NINO_HADLEY>>]], because they needed proof of identity from both sides. and I'm sat there thinking, who exactly am I reading this out to? turned out fine, but still. the registered agent is in [[L:austin|Austin]], and [[P:morrison_k|Karen]] handles the filings from [[L:dallas|Dallas]]. her number's [[PH:morrison_k_phone|(214) 555-0143]] if you ever need a US accountant, she's brilliant.

\[00:11:48\] **And how's the relationship with the marketplaces themselves?**

fees and fines. on [[d.ID.date|the 5th of March]] we got hit with a [[d.ID.amount|£600]] charge because a carton was two centimetres over the size limit. two centimetres. and you write to seller support and get a bot sending you templates. then there's inbound: you book a slot at the fulfilment centre, the lorry turns up and they turn it away, and you pay for the wait. [[d.ORG.platform.pub|Shopify]]'s nicer in that sense, it's our own store, but sales there are maybe a fifth of [[d.ORG.platform.pub|Amazon]]. so there's not much choice really. and the rules change every month, that's what kills you.

\[00:13:20\] **How do you find new products these days?**

analytics tools, mostly, you can see what's trending. plus [[P:lin|Wei]] sends me stuff he spots at trade fairs in [[L:shenzhen|Shenzhen]] and [[L:guangzhou|Guangzhou]], the [[d.ORG.event.pub|Canton Fair]] runs twice a year. he's going in April, I've already given him a list of about thirty items. [[P:ashdown|Robbo]] throws ideas in as well, more gadgets, but he's got a good eye. the hard part isn't finding a product, it's the timing. by the time stock lands two months later the niche is rammed. happened with water bottles, we were three weeks too late.

\[00:15:02\] **Do you have an office, or is everyone remote?**

there's a small office in [[L:leeds|Leeds]], it's really just a desk in a co-working space. and we've got a tiny admin office in [[L:harrogate|Harrogate]], one person doing invoices and paperwork, she's [[i|the only bookkeeper at our [[L:harrogate|Harrogate]] branch]], [[P:whitcombe|Sarah]] supervises her remotely. why [[L:harrogate|Harrogate]]? cheap rent and we found a good person, that's literally it. the warehouse, as I said, is outside [[L:bradford|Bradford]], that's [[P:ashdown|Rob's]] kingdom.

\[00:16:28\] **How does the team communicate? Which tools?**

WhatsApp, mostly. one big group, everyone dumps everything in there, photos, invoices, memes. [[P:ashdown|Rob]] posts photos of damaged deliveries. sometimes he messages at two in the morning, the lorry's in and half the boxes are crushed. [[P:whitcombe|Sarah]] prefers email, her work address is [[E:whitcombe_email|s.whitcombe@example.net]], she keeps all the invoices there, says things get lost in WhatsApp. and spreadsheets, obviously. you've got my number, but just in case, [[PH:hadley_phone|07700 900526]], that's the main one.

\[00:18:10\] **Were there disputes with contractors that really hurt the business?**

the freight thing was the big one. when that container went missing I genuinely thought we were done. loan to repay, no stock, bad reviews piling up. what saved us was a solicitor, [[P:fairbairn|J. T. Fairbairn]], from [[O:fairbairn_co|Fairbairn & Cole LLP]] in [[L:manchester|Manchester]]. he drafted the complaint and the letter before action, never went to court. we paid [[P:fairbairn|Fairbairn]] about [[d.ID.amount|£2,800]] and it was the best money we spent that year. he said [[O:pellow|Pellow]] had a whole stack of cases like ours, people just don't bother suing. and then there was the factory that sent spice racks without the screws, [[P:lin|Wei]] went over there himself and shouted at them for half a day.

\[00:19:50\] **Who makes the financial decisions?**

me. well, formally me, but no payment goes out without [[P:whitcombe|Sarah]] signing it off. on paper she's down as [[P:whitcombe|S. Whitcombe]], responsible for accounts, she's got bank access. I call her [[P:whitcombe|Ms Whitcombe]] when she's telling me off for overspending. for sourcing I've got a separate card, [[C:hadley_card|<<CARD_EN>>]], [[P:lin|Wei]] invoices against it and [[P:whitcombe|Sarah]] reconciles. anything big, like a new order over fifty grand, the three of us talk it through, me, [[P:whitcombe|Sarah]] and [[P:ashdown|Rob]], because he's the one who has to find space for it.

\[00:21:15\] **How much time a week goes on routine work you'd like to hand off?**

loads. twenty hours, maybe. reports, reconciliations, replying to reviews, rewriting listings. reviews are their own nightmare, dozens a day, and you have to answer or the rating drops. I tried giving it to a freelancer and he made it worse. now I do some myself and an AI tool does the rest. [[P:whitcombe|Sarah]] does the reports, but she's got a couple of other clients too, so she can't do everything.

\[00:22:35\] **Have you tried giving reports or spreadsheets to an AI tool?**

I have, yeah. and honestly that's where the paranoia kicks in. once I dropped our payroll sheet into a chatbot so it would work out the quarterly cost. full names, NI numbers, salaries, the lot. it did the maths beautifully, and then I sat there thinking, so where does all that live now? [[P:whitcombe|Sarah]] nearly killed me when she found out. she said, you know that's personal data, there are fines for that. [[P:ashdown|Rob Ashdown's]] got a worse story, he uploaded a supplier chat and the supplier's passport scan was in the attachments. nobody found out, but still.

\[00:24:20\] **What would you want a tool like that to do?**

drop in a spreadsheet or a document, it strips out the names, numbers, phone numbers, gives me something clean. then when the AI answers, it puts everything back, so I'm not sitting there working out who [[d.PERSON.label|PERSON_3]] is. and it has to run on my laptop, no cloud, otherwise what's the point. and it mustn't touch the numbers. if it swaps [[d.ID.amount|£210,000]] for some placeholder, I don't want it, the numbers are the whole point. same with dates, I need to know a shipment landed on [[d.ID.date|14 February]], not just at some point.

\[00:25:35\] **How do you hire people? Where do you find them?**

mostly through people we know, to be honest. we tried job boards, you get people who leave after a month. [[P:ashdown|Rob]] brought in both his packers, they'd worked together before. our newest listings person, [[P:pryor|Hannah Pryor]], came through a friend of a friend in [[L:manchester|Manchester]], she works remotely, does the listings and edits the photos. [[P:pryor|Hannah's]] got a real eye for design, listings sell maybe twenty per cent better after she's touched them. [[P:whitcombe|Sarah]] handles all the contracts, she knows what paperwork's needed, I don't go near it. and there's a one-month probation, if someone hasn't clicked by then they never will, in my experience.

\[00:27:05\] **Where do you keep staff documents?**

oh. this is the embarrassing bit. some of it's in a shared drive, in a folder called HR, some of it's on [[P:whitcombe|Sarah's]] laptop, and some of it's just in WhatsApp. when [[P:pryor|Hannah]] joined she dropped her passport scan straight into the group chat, and I think it's still sitting there, nobody ever deleted it. [[P:whitcombe|Sarah]] keeps the signed contracts and the originals. I know it's not how it should be, but when there's seven of you, you don't really think about policies. if someone came along and said, here's how you do it, here's the button, I'd do it. I just don't have the time to work it out myself.

\[00:28:30\] **Have you ever been audited?**

once, in [[d.ID.year|2024]]. we got a letter from [[d.ORG.gov.pub|HMRC]] asking us to explain some VAT figures on returns, something didn't add up. [[P:whitcombe|Sarah]] spent two days pulling the paperwork together, it all checked out in the end, but we still paid a [[d.ID.amount|£400]] penalty for filing something late. I gave her a raise afterwards, honestly. there was another time when stock arrived without the right labelling, but that was the marketplace fining us, not the government. lots of stress, lots of money, nothing to show for it.

\[00:30:00\] **Where do you see the business in a year?**

I want a second warehouse, probably near [[L:york|York]], closer to the north-east customers, freight works out cheaper. and my wife, [[P:hadley_e|Emma]], wants her own line, kids' stuff, she says I know nothing about it and she's right. [[P:hadley_e|Em's]] already looking at suppliers, through [[P:lin|Wei]] as well. if it works out it'll be a separate company in her name. [[P:ashdown|Robbo]] says I've lost my mind, but he says that about every idea I have and then copies it six months later. and I'd really like to get the paperwork sorted so I'm not embarrassed by it.

\[00:31:25\] **Who else do you know that might be interested?**

[[P:ashdown|Robbo]], definitely, he's mentioned it himself. there's also [[P:brennan|Niamh Brennan]] in [[L:belfast|Belfast]], she imports textiles, runs [[O:brennan_ltd|Brennan Linen Ltd]], about twenty staff, all the HR stuff lives in spreadsheets. her number's [[PH:brennan_phone|+44 28 9649 6012]], tell her I sent you. and I'd talk to [[i|the woman in [[O:northgate|Northgate]]'s purchasing team who won their supplier-of-the-year thing last year]], she's very organised. oh, and [[i|our forklift driver, the one who ran the London Marathon dressed as a dinosaur]], he keeps the shift rota in spreadsheets as well.

\[00:33:00\] **Last question. What would you pay for something like that?**

depends how well it works. if it really saves time and I'm not worried about the data, twenty, thirty quid a month, easy. with an [[d.ORG.platform.pub|Amazon]] integration so reports come out clean straight away, fifty would be fine. for comparison, I spend about [[d.ID.amount|$120]] a month on analytics and I think that's steep.

\[00:33:45\] **Thanks, [[P:hadley|Tom]]. Can I come back to you in a month with a prototype?**

yeah, of course. message me or call [[PH:hadley_phone|+44 7700 900526]], same number, I just say it differently every time. not Mondays though, Monday's inbound day, I'm at the warehouse with [[P:ashdown|Rob]] all day. and do get in touch with [[P:okafor|Grace Okafor]], seriously, she'll tell you loads about freight.
"""

# Belfast landline drama range is 028 9649 6000-6999, the rest use Ofcom's drama ranges
# (07700 900xxx, 020 7946 0xxx) and US 555-01xx.


TIMESTAMP_RE = re.compile(r"\\\[(\d\d:\d\d:\d\d)\\\]")


def build_text(file: str, template: str, subs: dict[str, str]):
    text, marks = render(template, subs)
    golden, decoys, indirect = classify(file, text, marks)
    for m in TIMESTAMP_RE.finditer(text):
        decoys.append(
            {
                "file": file,
                "text": m.group(1),
                "kind": "ID",
                "type": "timestamp",
                "start": m.start(1),
                "end": m.end(1),
            }
        )
    decoys.sort(key=lambda r: r["start"])
    return text, golden, decoys, indirect


# --------------------------------------------------------------------------
# Staff tables
# --------------------------------------------------------------------------

RU_SURNAMES = [
    ("Ковалев", "Ковалева"),
    ("Морозов", "Морозова"),
    ("Соколов", "Соколова"),
    ("Лебедев", "Лебедева"),
    ("Волков", "Волкова"),
    ("Зайцев", "Зайцева"),
    ("Белов", "Белова"),
    ("Орлов", "Орлова"),
    ("Комаров", "Комарова"),
    ("Федотов", "Федотова"),
    ("Тихонов", "Тихонова"),
    ("Голубев", "Голубева"),
    ("Ершов", "Ершова"),
    ("Савин", "Савина"),
    ("Широков", "Широкова"),
    ("Гаврилов", "Гаврилова"),
    ("Устинов", "Устинова"),
    ("Щербаков", "Щербакова"),
    ("Дорофеев", "Дорофеева"),
    ("Кондратьев", "Кондратьева"),
    ("Миронов", "Миронова"),
    ("Ларионов", "Ларионова"),
    ("Субботин", "Субботина"),
    ("Евсеев", "Евсеева"),
    ("Блинов", "Блинова"),
    ("Харитонов", "Харитонова"),
    ("Прохоров", "Прохорова"),
    ("Селезнев", "Селезнева"),
    ("Шилов", "Шилова"),
    ("Абрамов", "Абрамова"),
    ("Вешняков", "Вешнякова"),
    ("Гурьев", "Гурьева"),
    ("Демидов", "Демидова"),
    ("Жуков", "Жукова"),
    ("Исаев", "Исаева"),
    ("Калинин", "Калинина"),
    ("Нестеров", "Нестерова"),
    ("Осипов", "Осипова"),
    ("Пестов", "Пестова"),
    ("Рябов", "Рябова"),
    ("Котов", "Котова"),
    ("Панфилов", "Панфилова"),
]
RU_MALE = [
    "Алексей",
    "Дмитрий",
    "Андрей",
    "Игорь",
    "Павел",
    "Максим",
    "Никита",
    "Роман",
    "Виктор",
    "Егор",
    "Кирилл",
    "Владимир",
    "Степан",
    "Глеб",
]
RU_FEMALE = [
    "Анна",
    "Елена",
    "Ольга",
    "Татьяна",
    "Наталья",
    "Светлана",
    "Юлия",
    "Ксения",
    "Дарья",
    "Екатерина",
    "Алина",
    "Вера",
    "Полина",
    "Софья",
]
RU_FATHERS = [
    ("Андреевич", "Андреевна"),
    ("Сергеевич", "Сергеевна"),
    ("Викторович", "Викторовна"),
    ("Игоревич", "Игоревна"),
    ("Павлович", "Павловна"),
    ("Николаевич", "Николаевна"),
    ("Олегович", "Олеговна"),
    ("Юрьевич", "Юрьевна"),
    ("Михайлович", "Михайловна"),
    ("Александрович", "Александровна"),
]
RU_CITIES = ["Москва", "Москва", "Москва", "Подольск", "Тверь", "Казань", "Иваново"]
RU_DEPTS = [
    ("Склад", ["Кладовщик", "Комплектовщик", "Старший кладовщик"]),
    ("Закупки", ["Менеджер по закупкам", "Специалист ВЭД"]),
    ("Маркетплейсы", ["Менеджер маркетплейсов", "Контент-менеджер"]),
    ("Бухгалтерия", ["Бухгалтер", "Бухгалтер по первичке"]),
    ("Логистика", ["Логист", "Водитель-экспедитор"]),
]

RU_FIXED_ROLES = {
    "p01": ("Бухгалтерия", "Главный бухгалтер"),
    "p02": ("Закупки", "Руководитель закупок"),
    "p04": ("Склад", "Начальник склада"),
    "p05": ("Маркетплейсы", "Руководитель направления"),
    "p10": ("Бухгалтерия", "Бухгалтер"),
    "p13": ("Склад", "Старший кладовщик"),
}
EN_FIXED_ROLES = {
    "p01": ("Finance", "Finance Lead"),
    "p02": ("Warehouse", "Warehouse Manager"),
    "p03": ("Logistics", "Logistics Manager"),
    "p04": ("Finance (US)", "US Accountant"),
    "p05": ("Sourcing", "Sourcing Lead"),
    "p10": ("Finance", "Accounts Assistant"),
    "p13": ("Warehouse", "Shift Lead"),
}

EN_FIRST_M = [
    "James",
    "Oliver",
    "Harry",
    "George",
    "Callum",
    "Liam",
    "Ethan",
    "Owen",
    "Marcus",
    "Declan",
    "Rhys",
    "Nathan",
    "Aaron",
    "Felix",
]
EN_FIRST_F = [
    "Emily",
    "Chloe",
    "Hannah",
    "Megan",
    "Lucy",
    "Aisha",
    "Freya",
    "Imogen",
    "Priya",
    "Zoe",
    "Ellie",
    "Rosie",
    "Bethany",
    "Holly",
]
EN_SURNAMES = [
    "Whitcombe",
    "Ashdown",
    "Morrison",
    "Thornbury",
    "Callow",
    "Brennan",
    "Okafor",
    "Lindqvist",
    "Gallagher",
    "Ridley",
    "Hartigan",
    "Sutcliffe",
    "Kowalczyk",
    "Dunmore",
    "Pryce",
    "Ellery",
    "Marchetti",
    "Quayle",
    "Holloway",
    "Abernethy",
    "Fenwick",
    "Garside",
    "Haslam",
    "Ingram",
    "Jessop",
    "Kilbride",
    "Lomax",
    "Nuttall",
    "Ormerod",
    "Pickering",
    "Rawcliffe",
    "Spedding",
    "Tennant",
    "Varley",
    "Wainwright",
    "Yeadon",
    "Ackroyd",
    "Birtwistle",
    "Crowther",
    "Dransfield",
    "Entwistle",
    "Hepworth",
]
EN_CITIES_UK = ["Leeds", "Leeds", "Leeds", "Bradford", "Harrogate", "Manchester"]
EN_CITIES_US = ["Austin", "Dallas"]
EN_DEPTS = [
    ("Warehouse", ["Warehouse Operative", "Picker Packer", "Shift Lead"]),
    ("Sourcing", ["Sourcing Coordinator", "Import Specialist"]),
    ("Marketplaces", ["Marketplace Manager", "Listings Executive"]),
    ("Finance", ["Bookkeeper", "Accounts Assistant"]),
    ("Logistics", ["Logistics Coordinator", "Driver"]),
]


def _initials_ru(p: dict) -> str:
    return f"{p['last']} {p['first'][0]}.{p['patr'][0]}."


def make_ru_people(rng: random.Random, n: int) -> list[dict]:
    people: list[dict] = []
    fixed = [
        ("f", 0, "Мария", 0),  # Ковалева Мария Андреевна
        ("m", 1, "Олег", 2),  # Морозов Олег Викторович
        ("f", 1, "Анна", 1),  # Морозова Анна Сергеевна (namesake of the previous one)
        ("m", 3, "Дмитрий", 5),  # Лебедев Дмитрий Николаевич
        ("f", 2, "Елена", 4),  # Соколова Елена Павловна
    ]
    seen = set()
    for g, si, first, fi in fixed:
        last = RU_SURNAMES[si][0 if g == "m" else 1]
        patr = RU_FATHERS[fi][0 if g == "m" else 1]
        people.append({"g": g, "last": last, "first": first, "patr": patr})
        seen.add(last)
    pool = list(range(4, len(RU_SURNAMES)))
    rng.shuffle(pool)
    while len(people) < n:
        si = pool.pop()
        g = rng.choice("mf")
        last = RU_SURNAMES[si][0 if g == "m" else 1]
        first = rng.choice(RU_MALE if g == "m" else RU_FEMALE)
        patr = rng.choice(RU_FATHERS)[0 if g == "m" else 1]
        people.append({"g": g, "last": last, "first": first, "patr": patr})
    for i, p in enumerate(people, 1):
        p["id"] = f"p{i:02d}"
        p["full"] = f"{p['last']} {p['first']} {p['patr']}"
    return people


def build_ru_staff(path: Path, file: str):
    rng = random.Random(7001)
    people = make_ru_people(rng, 40)
    by = {p["id"]: p for p in people}
    golden, decoys, indirect = [], [], []

    def add_cell(ws, sheet: str, coord: str, template: str):
        text, marks = render(template)
        ws[coord] = text
        g, d, ind = classify(file, text, marks, {"where": f"{sheet}!{coord}"})
        golden.extend(g)
        decoys.extend(d)
        indirect.extend(ind)

    def mark_whole(sheet: str, coord: str, value: str, kind: str, entity: str):
        golden.append(
            {
                "file": file,
                "text": value,
                "kind": kind,
                "entity": entity,
                "where": f"{sheet}!{coord}",
                "start": 0,
                "end": len(value),
            }
        )

    # Comments: (row person id) -> template; references use entity ids of other rows.
    comments = {
        "p03": "Однофамилица [[P:p02|Морозова О.В.]], не родственники",
        "p04": "Замещает [[P:p01|Ковалеву М.А.]] на время отпуска",
        "p06": "Наставник: [[P:p02|Олег Морозов]]",
        "p07": "Переведена из команды [[P:p05|Соколовой]] [[d.ID.date|01.02.2026]]",
        "p09": "Испытательный срок до [[d.ID.date|01.06.2026]]",
        "p11": "Премию согласовала [[P:p01|Мария Андреевна]]",
        "p12": "Ведет кабинет [[d.ORG.platform.pub|Wildberries]] и [[d.ORG.platform.pub|Ozon]]",
        "p14": "По вопросам графика к [[P:p04|Лебедеву Д.Н.]]",
        "p17": "Оклад пересмотрен в [[d.ID.date|марте 2026]], +[[d.ID.percent|10%]]",
        "p19": "Сменщик: [[P:p08|<<P08_FIRST_LAST>>]]",
        "p22": "Работает удаленно, документы через [[P:p01|Ковалеву]]",
        "p25": "Совмещает с работой у [[P:p02|Морозова]] в закупках",
        "p28": "Декрет с [[d.ID.date|15.11.2025]]",
        "p31": "Курирует новичков вместе с [[P:p05|Еленой Соколовой]]",
        "p34": "Командировки в [[d.LOCATION.country.pub|Китай]] 2 раза в год",
        "p37": "Отчеты для [[P:p04|Дмитрия Николаевича]]",
    }
    p08 = by["p08"]
    subs_comment = {"P08_FIRST_LAST": f"{p08['first']} {p08['last']}"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    sheet = ws.title
    headers = [
        "ФИО",
        "Табельный номер",
        "СНИЛС",
        "Должность",
        "Отдел",
        "Город",
        "Оклад",
        "Премия %",
        "Годовой ФОТ",
        "Комментарий",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    tabs = rng.sample(range(100, 1000), len(people))
    snils_rng = random.Random(7002)
    used_snils = set()
    for r, p in enumerate(people, start=2):
        dept, roles = RU_DEPTS[rng.randrange(len(RU_DEPTS))]
        role = rng.choice(roles)
        if p["id"] in RU_FIXED_ROLES:
            dept, role = RU_FIXED_ROLES[p["id"]]
        tab = f"{tabs[r - 2]:05d}"
        s = snils(snils_rng)
        while s in used_snils:
            s = snils(snils_rng)
        used_snils.add(s)
        city = rng.choice(RU_CITIES)
        salary = rng.randrange(55, 180) * 1000
        bonus = rng.choice([0, 5, 10, 10, 15, 20])
        ws.cell(row=r, column=1, value=p["full"])
        mark_whole(sheet, f"A{r}", p["full"], "PERSON", p["id"])
        ws.cell(row=r, column=2, value=tab)
        mark_whole(sheet, f"B{r}", tab, "ID", f"tab_{p['id']}")
        ws.cell(row=r, column=3, value=s)
        mark_whole(sheet, f"C{r}", s, "SNILS", f"snils_{p['id']}")
        ws.cell(row=r, column=4, value=role)
        ws.cell(row=r, column=5, value=dept)
        ws.cell(row=r, column=6, value=city)
        mark_whole(sheet, f"F{r}", city, "LOCATION", f"city_{city}")
        ws.cell(row=r, column=7, value=salary)
        ws.cell(row=r, column=8, value=bonus)
        ws.cell(row=r, column=9, value=f"=G{r}*12*(1+H{r}/100)")
        if p["id"] in comments:
            add_cell(ws, sheet, f"J{r}", _fill(comments[p["id"]], subs_comment))
    for col, w in zip("ABCDEFGHIJ", [36, 16, 16, 26, 16, 12, 10, 10, 14, 50]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Руководители")
    s2 = ws2.title
    ws2.append(["Отдел", "Руководитель", "Заместитель", "Примечание"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    p01, p02, p04, p05 = by["p01"], by["p02"], by["p04"], by["p05"]
    p10, p13 = by["p10"], by["p13"]
    rows = [
        (
            "Бухгалтерия",
            f"[[P:p01|{_initials_ru(p01)}]]",
            f"[[P:p10|{p10['first']} {p10['last']}]]",
            "Подписывает платежки",
        ),
        (
            "Закупки",
            f"[[P:p02|{p02['first']} {p02['patr']} {p02['last']}]]",
            f"[[P:p03|{_initials_ru(by['p03'])}]]",
            f"Связь с поставщиками в [[d.LOCATION.country.pub|Китае]]",
        ),
        (
            "Склад",
            f"[[P:p04|{p04['last']} {p04['first'][0]}. {p04['patr'][0]}.]]",
            f"[[P:p13|{p13['last']} {p13['first']}]]",
            "",
        ),
        (
            "Маркетплейсы",
            f"[[P:p05|{p05['first'][0]}.{p05['patr'][0]}. {p05['last']}]]",
            "",
            f"Временно также ведет логистику, пока [[P:p02|{p02['last']}]] в командировке",
        ),
        (
            "Логистика",
            f"[[P:p04|{p04['first']} {p04['last']}]]",
            "",
            "Вакансия заместителя открыта",
        ),
    ]
    for r, (dept, boss, dep, note) in enumerate(rows, start=2):
        ws2.cell(row=r, column=1, value=dept)
        add_cell(ws2, s2, f"B{r}", boss)
        if dep:
            add_cell(ws2, s2, f"C{r}", dep)
        if note:
            add_cell(ws2, s2, f"D{r}", note)
    for col, w in zip("ABCD", [16, 34, 30, 60]):
        ws2.column_dimensions[col].width = w

    _save(wb, path)
    return golden, decoys, indirect


def _fill(template: str, subs: dict[str, str]) -> str:
    for k, v in subs.items():
        template = template.replace(f"<<{k}>>", v)
    return template


def build_en_staff(path: Path, file: str):
    rng = random.Random(8001)
    golden, decoys, indirect = [], [], []
    fixed = [
        ("f", "Sarah", "Whitcombe"),
        ("m", "Rob", "Ashdown"),
        ("m", "Dave", "Morrison"),
        ("f", "Karen", "Morrison"),  # namesake, unrelated
        ("f", "Grace", "Okafor"),
    ]
    people = [{"g": g, "first": f, "last": l} for g, f, l in fixed]
    pool = [
        s
        for s in EN_SURNAMES
        if s not in {"Whitcombe", "Ashdown", "Morrison", "Okafor"}
    ]
    rng.shuffle(pool)
    while len(people) < 40:
        g = rng.choice("mf")
        people.append(
            {
                "g": g,
                "first": rng.choice(EN_FIRST_M if g == "m" else EN_FIRST_F),
                "last": pool.pop(),
            }
        )
    for i, p in enumerate(people, 1):
        p["id"] = f"p{i:02d}"
        p["full"] = f"{p['first']} {p['last']}"
    by = {p["id"]: p for p in people}

    def add_cell(ws, sheet, coord, template):
        text, marks = render(template)
        ws[coord] = text
        g, d, ind = classify(file, text, marks, {"where": f"{sheet}!{coord}"})
        golden.extend(g)
        decoys.extend(d)
        indirect.extend(ind)

    def mark_whole(sheet, coord, value, kind, entity):
        golden.append(
            {
                "file": file,
                "text": value,
                "kind": kind,
                "entity": entity,
                "where": f"{sheet}!{coord}",
                "start": 0,
                "end": len(value),
            }
        )

    notes = {
        "p03": "No relation to [[P:p04|Karen Morrison]]",
        "p06": "Covering for [[P:p01|Sarah Whitcombe]] during her leave",
        "p07": "Reports to [[P:p02|Rob Ashdown]]",
        "p09": "Probation ends [[d.ID.date|1 June 2026]]",
        "p11": "Bonus signed off by [[P:p01|S. Whitcombe]]",
        "p12": "Runs the [[d.ORG.platform.pub|Amazon]] and [[d.ORG.platform.pub|Shopify]] accounts",
        "p14": "Shift questions go to [[P:p02|Robbo]]",
        "p17": "Pay reviewed in [[d.ID.date|March 2026]], +[[d.ID.percent|8%]]",
        "p20": "[[P:p05|Grace's]] backup on freight bookings",
        "p23": "Works remotely, paperwork via [[P:p01|Sarah]]",
        "p26": "Shadowing [[P:p03|Dave]] on customs entries",
        "p29": "Parental leave from [[d.ID.date|15/11/2025]]",
        "p33": "Sourcing trips to [[d.LOCATION.country.pub|China]] twice a year",
        "p36": "Prepares reports for [[P:p04|Ms Morrison]]",
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    sheet = ws.title
    ws.append(
        [
            "Full name",
            "Employee ID",
            "Email",
            "Phone",
            "Department",
            "City",
            "Salary",
            "Bonus %",
            "Annual cost",
            "Notes",
        ]
    )
    for c in ws[1]:
        c.font = Font(bold=True)
    uk_tail = rng.sample(range(100, 1000), 40)
    us_tail = rng.sample(range(100, 200), 40)
    emp_ids = rng.sample(range(1000, 2000), 40)
    for r, p in enumerate(people, start=2):
        dept, roles = EN_DEPTS[rng.randrange(len(EN_DEPTS))]
        role = rng.choice(roles)
        if p["id"] in EN_FIXED_ROLES:
            dept, role = EN_FIXED_ROLES[p["id"]]
        us = p["id"] == "p04" or (p["id"] not in EN_FIXED_ROLES and rng.random() < 0.2)
        city = "Dallas" if p["id"] == "p04" else rng.choice(EN_CITIES_US if us else EN_CITIES_UK)
        phone = f"(512) 555-0{us_tail[r - 2]}" if us else f"07700 900{uk_tail[r - 2]}"
        if us and city == "Dallas":
            phone = "(214) " + phone[6:]
        email = f"{p['first'].lower()}.{p['last'].lower()}@example.com"
        emp = f"EMP-{emp_ids[r - 2]}"
        salary = (
            (rng.randrange(220, 600) * 100)
            if not us
            else (rng.randrange(380, 900) * 100)
        )
        bonus = rng.choice([0, 5, 10, 10, 15])
        vals = [
            p["full"],
            emp,
            email,
            phone,
            dept,
            city,
            salary,
            bonus,
            f"=G{r}*(1+H{r}/100)",
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        mark_whole(sheet, f"A{r}", p["full"], "PERSON", p["id"])
        mark_whole(sheet, f"B{r}", emp, "ID", f"emp_{p['id']}")
        mark_whole(sheet, f"C{r}", email, "EMAIL", f"email_{p['id']}")
        mark_whole(sheet, f"D{r}", phone, "PHONE", f"phone_{p['id']}")
        mark_whole(sheet, f"F{r}", city, "LOCATION", f"city_{city}")
        if p["id"] in notes:
            add_cell(ws, sheet, f"J{r}", notes[p["id"]])
    for col, w in zip("ABCDEFGHIJ", [24, 12, 34, 16, 14, 12, 10, 10, 12, 50]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Managers")
    s2 = ws2.title
    ws2.append(["Department", "Manager", "Deputy", "Note"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    p10, p13 = by["p10"], by["p13"]
    rows = [
        (
            "Finance",
            "[[P:p01|S. Whitcombe]]",
            f"[[P:p10|{p10['first'][0]}. {p10['last']}]]",
            "Signs off all payments",
        ),
        (
            "Warehouse",
            "[[P:p02|Ashdown, Rob]]",
            f"[[P:p13|{p13['last']}, {p13['first']}]]",
            "",
        ),
        (
            "Logistics",
            "[[P:p03|D. Morrison]]",
            "",
            "Not the same person as [[P:p04|K. Morrison]] in Finance (US)",
        ),
        (
            "Finance (US)",
            "[[P:p04|Karen Morrison]]",
            "",
            "Based in [[L:city_Dallas|Dallas]]",
        ),
        (
            "Sourcing",
            "[[P:p05|Grace Okafor]]",
            "",
            "Temporary, until [[P:p03|Dave]] is back",
        ),
    ]
    for r, (dept, boss, dep, note) in enumerate(rows, start=2):
        ws2.cell(row=r, column=1, value=dept)
        add_cell(ws2, s2, f"B{r}", boss)
        if dep:
            add_cell(ws2, s2, f"C{r}", dep)
        if note:
            add_cell(ws2, s2, f"D{r}", note)
    for col, w in zip("ABCD", [16, 24, 24, 60]):
        ws2.column_dimensions[col].width = w

    _save(wb, path)
    return golden, decoys, indirect


def _save(wb: Workbook, path: Path) -> None:
    fixed = datetime(2026, 1, 1, 0, 0, 0)
    wb.properties.creator = "interview-anonymizer fixtures"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _normalize_zip(path)


def _normalize_zip(path: Path) -> None:
    """Rewrite the xlsx container with fixed timestamps so reruns give identical bytes."""
    import zipfile

    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    stamp = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = stamp.sub(rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)


# --------------------------------------------------------------------------
# Verification
# --------------------------------------------------------------------------


def _word_re(surface: str) -> re.Pattern:
    return re.compile(r"(?<!\w)" + re.escape(surface) + r"(?!\w)", re.IGNORECASE)


def verify_text(file: str, text: str, golden, decoys, indirect) -> list[str]:
    errs = []
    spans = sorted((g["start"], g["end"], g) for g in golden)
    for s, e, g in spans:
        if text[s:e] != g["text"]:
            errs.append(
                f"{file}: span {s}:{e} is {text[s:e]!r}, golden says {g['text']!r}"
            )
        if g["text"] != g["text"].strip() or not g["text"]:
            errs.append(f"{file}: golden surface has stray whitespace: {g['text']!r}")
    for (s1, e1, g1), (s2, e2, g2) in zip(spans, spans[1:]):
        if s2 < e1:
            errs.append(
                f"{file}: golden mentions overlap: {g1['text']!r} / {g2['text']!r}"
            )
    for d in decoys:
        if text[d["start"] : d["end"]] != d["text"]:
            errs.append(f"{file}: decoy span mismatch {d['text']!r}")
        for s, e, g in spans:
            if s < d["end"] and d["start"] < e:
                errs.append(
                    f"{file}: decoy {d['text']!r} overlaps golden {g['text']!r}"
                )
    for i in indirect:
        if text[i["start"] : i["end"]] != i["text"]:
            errs.append(f"{file}: indirect span mismatch {i['text']!r}")
    # every occurrence of a golden surface must be covered by a golden span
    for surface in {g["text"] for g in golden}:
        for m in _word_re(surface).finditer(text):
            if not any(s <= m.start() and m.end() <= e for s, e, _ in spans):
                errs.append(
                    f"{file}: unmarked occurrence of golden surface {surface!r} at {m.start()}"
                )
    return errs


def verify_xlsx(file: str, path: Path, golden, decoys, indirect) -> list[str]:
    errs = []
    wb = load_workbook(path)
    by_cell: dict[str, list[dict]] = {}
    for g in golden:
        by_cell.setdefault(g["where"], []).append(g)
    for rec in [*golden, *decoys, *indirect]:
        sheet, coord = rec["where"].split("!")
        v = wb[sheet][coord].value
        v = "" if v is None else str(v)
        if v[rec["start"] : rec["end"]] != rec["text"]:
            errs.append(
                f"{file}: {rec['where']} holds {v!r}, expected {rec['text']!r} at {rec['start']}"
            )
    surfaces = {g["text"] for g in golden}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or c.value.startswith("="):
                    continue
                where = f"{ws.title}!{c.coordinate}"
                cover = [(g["start"], g["end"]) for g in by_cell.get(where, [])]
                for surface in surfaces:
                    for m in _word_re(surface).finditer(c.value):
                        if not any(s <= m.start() and m.end() <= e for s, e in cover):
                            errs.append(
                                f"{file}: unmarked {surface!r} in {where}: {c.value!r}"
                            )
    return errs


def verify_ids(golden) -> list[str]:
    errs = []
    for g in golden:
        if g["kind"] not in VALID_KINDS:
            errs.append(f"invalid kind {g['kind']}")
        if g["kind"] == "SNILS" and not snils_valid(g["text"]):
            errs.append(f"bad SNILS checksum {g['text']}")
        if g["kind"] == "CARD" and not luhn_valid(g["text"]):
            errs.append(f"bad card checksum {g['text']}")
    return errs


# --------------------------------------------------------------------------


def main() -> int:
    rng = random.Random(20261001)
    subs = {
        "SNILS_GUSEV": snils(rng),
        "INN_GUSEV": inn12(rng),
        "INN_BAIKAL": inn10(rng),
        "CARD_RU": luhn_card(rng, "400012"),
        "SSN_HADLEY": ssn(rng),
        "NINO_HADLEY": nino(rng),
        "CARD_EN": luhn_card(rng, "400034"),
    }

    golden: list[dict] = []
    decoys: list[dict] = []
    indirect: list[dict] = []
    files: dict[str, dict] = {}
    errors: list[str] = []

    for file, lang, tpl in [
        ("ru/interview.md", "ru", RU_INTERVIEW),
        ("en/interview.md", "en", EN_INTERVIEW),
    ]:
        text, g, d, i = build_text(file, tpl, subs)
        path = FIX / file
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        reread = path.read_text(encoding="utf-8")
        errors += verify_text(file, reread, g, d, i)
        golden += g
        decoys += d
        indirect += i
        files[file] = {
            "type": "text",
            "lang": lang,
            "words": len(re.findall(r"\w+", text)),
        }

    for file, lang, builder in [
        ("ru/staff.xlsx", "ru", build_ru_staff),
        ("en/staff.xlsx", "en", build_en_staff),
    ]:
        path = FIX / file
        g, d, i = builder(path, file)
        errors += verify_xlsx(file, path, g, d, i)
        golden += g
        decoys += d
        indirect += i
        files[file] = {"type": "xlsx", "lang": lang}

    errors += verify_ids(golden)
    if errors:
        for e in errors:
            print("ERROR:", e, file=sys.stderr)
        raise SystemExit(f"{len(errors)} fixture errors, golden.json not written")

    out = {
        "note": "Synthetic reference set. All people, companies, numbers are invented. "
        "start/end are character offsets into the file text (for xlsx: into the cell value at `where`).",
        "files": files,
        "mentions": golden,
        "decoys": decoys,
        "indirect": indirect,
    }
    (FIX / "golden.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )

    # stats
    for file, meta in files.items():
        g = [x for x in golden if x["file"] == file]
        kinds: dict[str, int] = {}
        for x in g:
            kinds[x["kind"]] = kinds.get(x["kind"], 0) + 1
        nd = sum(1 for x in decoys if x["file"] == file)
        ni = sum(1 for x in indirect if x["file"] == file)
        ents = len({x["entity"] for x in g})
        extra = f", {meta['words']} words" if "words" in meta else ""
        print(
            f"{file}: {len(g)} mentions, {ents} entities, {nd} decoys, {ni} indirect{extra}"
        )
        print("   " + ", ".join(f"{k} {v}" for k, v in sorted(kinds.items())))
    print(f"wrote {FIX / 'golden.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
