"""Tables (.xlsx, .csv): column classification and cell-by-cell anonymization.

scan_table() finds the header row of every sheet, classifies columns
(by header synonyms first, by values second) and runs the detectors over the
remaining free-text cells. apply_table() writes a copy where classified
columns hold vault labels and free-text cells went through replace.apply().

Numbers, dates, formulas and header cells are never changed, with one
exception: a column whose header names an identifier (phone, INN, SNILS,
personnel number...) is replaced even when its values are stored as numbers,
because those numbers are the sensitive value itself, not an amount.

The detectors (anon.detect) and the replacer (anon.replace) are imported
lazily through _detector() and _replacer(), so this module imports without
them and tests can swap them for fakes.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from . import lang as _lang
from . import names
from .model import Kind, Merged
from .vault import Entity, Vault, fold

HEADER_SCAN_ROWS = 10  # header row is searched among the first rows only
VALUE_SHARE = 0.6  # share of cells one kind must cover to classify a column by values
MIN_VALUE_CELLS = 2  # fewer text cells than this: no value-based classification

XLSX_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}

# Kinds whose values are the identifier itself, so numeric storage is no excuse to keep them.
IDENTIFIER_KINDS = {
    Kind.PHONE,
    Kind.INN,
    Kind.SNILS,
    Kind.PASSPORT,
    Kind.SSN,
    Kind.NINO,
    Kind.ID,
    Kind.CARD,
}

PART_LAST = "last"
PART_FIRST = "first"
PART_MIDDLE = "middle"
_PART_ORDER = {
    "ru": (PART_LAST, PART_FIRST, PART_MIDDLE),  # Фамилия Имя Отчество
    "en": (PART_FIRST, PART_MIDDLE, PART_LAST),  # First Middle Last
}


# ---- detector hooks (lazy, swappable in tests) ---------------------------


def _detector() -> Callable[..., list[Merged]]:
    from .detect import find_all

    return find_all


def _replacer() -> Callable[[str, list[Merged], Vault, str], str]:
    from .replace import apply

    return apply


# ---- public data ---------------------------------------------------------


@dataclass
class ColumnInfo:
    index: int  # 1-based column number
    letter: str  # "A", "B", ...
    header: str  # header text as shown (both levels for a two-row header)
    kind: Kind | None = None
    part: str | None = None  # "last" | "first" | "middle" for split name columns
    reason: str = ""  # "header" | "values" | "" when not classified
    numeric: bool = False  # most data cells are numbers
    dates: bool = False  # most data cells are dates
    group: int = 0  # name-part group: parts of one person share a group number


@dataclass
class TableItem:
    """One value for the review step; Merged-like plus where it sits."""

    kind: Kind
    text: str
    sheet: str
    cell: str  # "B2"
    context: str
    sources: list[str] = field(default_factory=list)
    start: int = 0  # offsets inside the cell text
    end: int = 0
    normal: str = ""


@dataclass
class SheetScan:
    name: str
    header_row: int | None  # 1-based, None when no header row was found
    header_rows: list[int]  # one or two rows (two for a merged two-level header)
    data_start: int  # first data row, 1-based
    columns: dict[int, ColumnInfo]  # every column seen, classified or not
    findings: dict[tuple[int, int], list[Merged]]  # free-text cell (row, col) -> spans
    cell_texts: dict[tuple[int, int], str]  # the scanned text of each free-text cell
    items: list[TableItem] = field(default_factory=list)

    def classified(self) -> list[ColumnInfo]:
        return [c for c in self.columns.values() if c.kind is not None]


@dataclass
class TableScan:
    path: Path
    format: str  # "xlsx" | "csv"
    sheets: list[SheetScan]
    delimiter: str = ","  # csv only
    encoding: str = "utf-8"  # csv only, as read

    def items(self) -> list[TableItem]:
        return [it for s in self.sheets for it in s.items]

    def sheet(self, name: str) -> SheetScan | None:
        for s in self.sheets:
            if s.name == name:
                return s
        return None

    def headers_text(self) -> str:
        return " ".join(
            c.header for s in self.sheets for c in s.columns.values() if c.header
        )


# ---- cell values ---------------------------------------------------------


class _FormulaType:
    def __repr__(self) -> str:
        return "<formula>"


FORMULA = _FormulaType()

_NUM = r"[+-]?(?:\d{1,3}(?:[   ,]\d{3})+|\d+)(?:[.,]\d+)?"
_NUM_RE = re.compile(
    rf"^\(?\s*[$€£₽]?\s*{_NUM}\s*(?:%|[$€£₽]|руб\.?|р\.|usd|eur|rub|gbp)?\s*\)?$",
    re.IGNORECASE,
)
_DATE_RES = [
    re.compile(r"^\d{1,2}[./-]\d{1,2}[./-]\d{2,4}(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?$"),
    re.compile(r"^\d{4}[./-]\d{1,2}[./-]\d{1,2}(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?$"),
    re.compile(r"^\d{1,2}:\d{2}(?::\d{2})?$"),
]
_PLACEHOLDERS = {"н/д", "нд", "n/a", "na", "нет", "none", "null", "nan", "-", "—", "–"}


def _cell_type(v: Any) -> str:
    """ "empty" | "text" | "number" | "date" | "bool" | "formula" | "other"."""
    if v is None or v is FORMULA:
        return "empty" if v is None else "formula"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float, Decimal)):
        return "number"
    if isinstance(v, (datetime, date, time, timedelta)):
        return "date"
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return "empty"
        if s.startswith("="):
            return "formula"
        if _NUM_RE.match(s):
            return "number"
        if any(r.match(s) for r in _DATE_RES):
            return "date"
        return "text"
    return "other"


def _is_placeholder(s: str) -> bool:
    s = s.strip().lower()
    return s in _PLACEHOLDERS or not re.search(r"[^\W_]", s)


def _surface(v: Any) -> str:
    """String form of a cell value used as the vault surface."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, Decimal) and v == v.to_integral_value():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def _col_letter(c: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(c)


# ---- sheet access (xlsx and csv behind one shape) ------------------------


@dataclass
class _Sheet:
    name: str
    grid: dict[tuple[int, int], Any]  # (row, col) -> value, FORMULA for formulas
    max_row: int
    max_col: int
    merged: list[tuple[int, int, int, int]]  # (min_row, min_col, max_row, max_col)
    setter: Callable[[int, int, Any], None] | None = None

    def get(self, r: int, c: int) -> Any:
        return self.grid.get((r, c))

    def set(self, r: int, c: int, v: Any) -> None:
        self.grid[(r, c)] = v
        if self.setter is not None:
            self.setter(r, c, v)

    def origin(self, r: int, c: int) -> tuple[int, int]:
        """Top-left cell of the merged range covering (r, c), or (r, c) itself."""
        for r0, c0, r1, c1 in self.merged:
            if r0 <= r <= r1 and c0 <= c <= c1:
                return r0, c0
        return r, c


def _xlsx_sheets(wb) -> list[_Sheet]:
    out = []
    for ws in wb.worksheets:
        grid: dict[tuple[int, int], Any] = {}
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                if cell.data_type == "f" or not isinstance(
                    v,
                    (str, int, float, Decimal, bool, datetime, date, time, timedelta),
                ):
                    # Formulas (plain, array, data table) and exotic objects stay untouched.
                    grid[(cell.row, cell.column)] = (
                        FORMULA if cell.data_type == "f" else v
                    )
                else:
                    grid[(cell.row, cell.column)] = v
        merged = [
            (m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges
        ]

        def setter(r: int, c: int, v: Any, ws=ws) -> None:
            ws.cell(row=r, column=c).value = v

        out.append(
            _Sheet(ws.title, grid, ws.max_row or 0, ws.max_column or 0, merged, setter)
        )
    return out


def _load_xlsx(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")


def _decode(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8"), "utf-8-sig"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1251"), "cp1251"


def _sniff(text: str) -> str:
    lines = [ln for ln in text.splitlines() if ln.strip()][:20]
    if not lines:
        return ","
    sample = "\n".join(lines)
    best, best_score = ",", (-1.0, 0)
    for d in (",", ";", "\t"):
        widths = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d)]
        if not widths:
            continue
        mode = max(set(widths), key=widths.count)
        if mode < 2:
            continue
        score = (widths.count(mode) / len(widths), mode)
        if score > best_score:
            best, best_score = d, score
    return best


def _read_csv(path: Path) -> tuple[list[list[str]], str, str]:
    text, enc = _decode(path.read_bytes())
    delim = _sniff(text)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    return rows, delim, enc


def _csv_sheet(name: str, rows: list[list[str]]) -> _Sheet:
    grid: dict[tuple[int, int], Any] = {}
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            if v != "":
                grid[(r, c)] = v

    def setter(r: int, c: int, v: Any) -> None:
        rows[r - 1][c - 1] = v

    return _Sheet(
        name, grid, len(rows), max((len(r) for r in rows), default=0), [], setter
    )


# ---- header synonyms -----------------------------------------------------

# Non-person kinds: matched anywhere in the normalized header, in this order
# (ID before PERSON words: "Employee ID"; EMAIL before LOCATION: "Email address").
_KIND_PATTERNS: list[tuple[Kind, re.Pattern[str]]] = [
    (
        Kind.ID,
        re.compile(
            r"\bтабельн\w*|\bтаб\b"
            r"|\b(?:employee|staff|personnel|emp|worker)\s+(?:id|no|number|num|№)\b"
            r"|\bemployee\s*id\b|\bpersonnel\s+number\b"
        ),
    ),
    (Kind.SNILS, re.compile(r"снилс")),
    (Kind.INN, re.compile(r"\bинн\b")),
    (Kind.PASSPORT, re.compile(r"паспорт|passport")),
    (Kind.SSN, re.compile(r"\bssn\b|social\s+security")),
    (Kind.NINO, re.compile(r"\bnino\b|\bni\s+(?:number|no)\b|national\s+insurance")),
    (
        Kind.EMAIL,
        re.compile(
            r"\be\s*mail|\bemail|\bmail\b"
            r"|\bпочт(?:а|ы|е|у|ой)\b|\bэл\w*\s+почт|\bemail"
        ),
    ),
    (Kind.PHONE, re.compile(r"телефон|\bтел\b|phone|\bmobile\b|\bмоб\w*")),
    (
        Kind.LOCATION,
        re.compile(
            r"адрес|address|\bгород\b|\bcity\b|\btown\b|населенный\s+пункт"
            r"|место\s+(?:жительства|проживания)"
        ),
    ),
]

_EN_PARTS: list[tuple[str, re.Pattern[str]]] = [
    (PART_LAST, re.compile(r"\b(?:last|family)\s+name\b|\bsurname\b|\blastname\b")),
    (PART_FIRST, re.compile(r"\b(?:first|given)\s+name\b|\bforename\b|\bfirstname\b")),
    (PART_MIDDLE, re.compile(r"\bmiddle\s+name\b|\bpatronymic\b")),
]

# Russian: the head noun comes first ("Руководитель отдела" vs "Отдел сотрудника").
_RU_PERSON = (
    "сотрудник",
    "работник",
    "руководител",
    "менеджер",
    "ответственн",
    "исполнител",
    "заместител",
    "директор",
    "начальник",
    "куратор",
    "наставник",
    "автор",
    "респондент",
    "модератор",
    "интервьюер",
    "рекрутер",
)
_RU_STOP = (
    "отдел",
    "подраздел",
    "департамент",
    "должност",
    "наименован",
    "назван",
    "количеств",
    "число",
    "кол",
    "дата",
    "сумм",
    "код",
    "номер",
    "оклад",
    "зарплат",
    "статус",
    "тип",
    "вид",
    "файл",
    "грейд",
    "ставк",
)
# "Имя файла", "Имя пользователя": the object after "имя" is not a person.
_RU_NAME_OBJECTS = (
    "файл",
    "пользовател",
    "лист",
    "пол",
    "переменн",
    "хост",
    "домен",
    "проект",
    "компани",
    "организаци",
    "продукт",
    "товар",
    "сервер",
    "баз",
    "таблиц",
    "документ",
    "бренд",
    "аккаунт",
    "учетн",
    "отдел",
)

# English: the head noun comes last ("Department manager" vs "Employee department").
_EN_PERSON = {
    "employee",
    "employees",
    "manager",
    "owner",
    "contact",
    "contacts",
    "assignee",
    "supervisor",
    "person",
    "worker",
    "deputy",
    "director",
    "head",
    "lead",
    "mentor",
    "author",
    "respondent",
    "interviewee",
    "interviewer",
    "moderator",
    "recruiter",
    "staff",
}
_EN_STOP = {
    "department",
    "dept",
    "division",
    "team",
    "position",
    "title",
    "role",
    "job",
    "company",
    "organization",
    "organisation",
    "org",
    "product",
    "project",
    "file",
    "sheet",
    "field",
    "date",
    "count",
    "number",
    "no",
    "id",
    "salary",
    "status",
    "type",
    "code",
    "host",
    "domain",
    "username",
    "user",
    "brand",
    "item",
    "account",
    "task",
    "city",
    "country",
    "grade",
    "level",
    "display",
    "event",
    "group",
    "category",
}


def _norm_header(s: str) -> str:
    s = str(s).lower().replace("ё", "е")
    return " ".join(re.findall(r"[a-zа-я0-9№]+", s))


def classify_header(header: str) -> tuple[Kind | None, str | None]:
    """(kind, name part) for a header text; (None, None) when it says nothing."""
    h = _norm_header(header)
    if not h:
        return None, None
    for kind, pat in _KIND_PATTERNS:
        if pat.search(h):
            return kind, None
    for part, pat in _EN_PARTS:
        if pat.search(h):
            return Kind.PERSON, part

    toks = h.split()
    # "Ф.И.О." -> "ф и о"
    if re.search(r"\bф и о\b", h):
        return Kind.PERSON, None

    ru = [t for t in toks if re.match(r"[а-я]", t)]
    for i, t in enumerate(ru):
        if t == "фио":
            return Kind.PERSON, None
        if t.startswith("фамили"):
            return Kind.PERSON, PART_LAST
        if t.startswith("отчеств"):
            return Kind.PERSON, PART_MIDDLE
        if t == "имя":
            nxt = ru[i + 1] if i + 1 < len(ru) else ""
            if nxt and nxt.startswith(_RU_NAME_OBJECTS):
                return None, None
            return Kind.PERSON, PART_FIRST
        if t.startswith("ответственност"):
            return None, None
        if t.startswith(_RU_PERSON):
            return Kind.PERSON, None
        if t.startswith("контактн") and "лицо" in ru[i + 1 : i + 2]:
            return Kind.PERSON, None
        if t.startswith(_RU_STOP):
            break

    en = [t for t in toks if re.match(r"[a-z]", t)]
    for i in range(len(en) - 1, -1, -1):
        t = en[i]
        if t == "name" or t == "fullname":
            prev = en[i - 1] if i > 0 else ""
            if prev in _EN_STOP:
                return None, None
            return Kind.PERSON, None
        if t in _EN_PERSON:
            return Kind.PERSON, None
        if t in _EN_STOP:
            break
    return None, None


# ---- scanning ------------------------------------------------------------


def _header_text(sh: _Sheet, r: int, c: int) -> str:
    v = sh.get(*sh.origin(r, c))
    if v is None or v is FORMULA:
        return ""
    return _surface(v)


def _find_header(sh: _Sheet) -> list[int]:
    for r in range(1, min(HEADER_SCAN_ROWS, sh.max_row) + 1):
        n = sum(
            1 for c in range(1, sh.max_col + 1) if _cell_type(sh.get(r, c)) == "text"
        )
        if n >= 2:
            rows = [r]
            if r + 1 <= sh.max_row and _two_level(sh, r):
                rows.append(r + 1)
            return rows
    return []


def _two_level(sh: _Sheet, r: int) -> bool:
    """Header row r has merged cells with sub-headers in row r + 1."""
    for r0, c0, r1, c1 in sh.merged:
        if r0 != r:
            continue
        if r1 > r:  # a header cell merged down over the sub-header row
            return True
        if c1 > c0 and any(
            _cell_type(sh.get(r + 1, c)) == "text" for c in range(c0, c1 + 1)
        ):
            return True
    return False


def _full_cover(text: str, spans: list[Merged]) -> Kind | None:
    """Kind of a single finding covering the whole (stripped) cell text."""
    lead = len(text) - len(text.lstrip())
    end = len(text.rstrip())
    for s in spans:
        if s.start <= lead and s.end >= end:
            return s.kind
    return None


def _row_context(sh: _Sheet, r: int, c: int) -> str:
    for cc in range(1, sh.max_col + 1):
        if cc == c:
            continue
        v = sh.get(r, cc)
        if _cell_type(v) == "text":
            return _surface(v)
    return ""


def _context(header: str, row_ctx: str) -> str:
    return " | ".join(x for x in (header, row_ctx) if x)


def _window(text: str, start: int, end: int, width: int = 40) -> str:
    a, b = max(0, start - width), min(len(text), end + width)
    return (
        ("..." if a else "")
        + text[a:b].replace("\n", " ")
        + ("..." if b < len(text) else "")
    )


def _script_lang(s: str) -> str:
    return "ru" if re.search(r"[А-Яа-яЁё]", s) else "en"


def _person_from_parts(parts: dict[str, str]) -> str:
    order = _PART_ORDER[_script_lang(" ".join(parts.values()))]
    return " ".join(parts[p] for p in order if parts.get(p))


def _scan_sheet(sh: _Sheet, find_all: Callable[..., list[Merged]]) -> SheetScan:
    header_rows = _find_header(sh)
    data_start = header_rows[-1] + 1 if header_rows else 1
    data_rows = range(data_start, sh.max_row + 1)

    columns: dict[int, ColumnInfo] = {}
    findings: dict[tuple[int, int], list[Merged]] = {}
    cell_texts: dict[tuple[int, int], str] = {}
    items: list[TableItem] = []
    cache: dict[str, list[Merged]] = {}

    def detect(text: str) -> list[Merged]:
        if text not in cache:
            cache[text] = list(find_all(text))
        return cache[text]

    for c in range(1, sh.max_col + 1):
        texts = [_header_text(sh, r, c) for r in header_rows]
        shown = " / ".join(dict.fromkeys(t for t in texts if t))
        own = next(
            (t for t in reversed(texts) if t), ""
        )  # the lowest header level wins
        col = ColumnInfo(index=c, letter=_col_letter(c), header=shown)
        columns[c] = col

        counts = {"text": 0, "number": 0, "date": 0}
        for r in data_rows:
            t = _cell_type(sh.get(r, c))
            if t in counts:
                counts[t] += 1
        col.dates = (
            counts["date"] > counts["text"] and counts["date"] >= counts["number"]
        )
        col.numeric = not col.dates and counts["number"] > counts["text"]

        kind, part = classify_header(own) if own else (None, None)
        if (
            kind is not None
            and not col.dates
            and (not col.numeric or kind in IDENTIFIER_KINDS)
        ):
            col.kind, col.part, col.reason = kind, part, "header"
            continue

        if col.numeric or col.dates:
            continue
        text_cells = [
            (r, sh.get(r, c)) for r in data_rows if _cell_type(sh.get(r, c)) == "text"
        ]
        if len(text_cells) >= MIN_VALUE_CELLS:
            tally: dict[Kind, int] = {}
            for _, v in text_cells:
                k = _full_cover(v, detect(v))
                if k is not None:
                    tally[k] = tally.get(k, 0) + 1
            if tally:
                best = max(tally, key=lambda k: tally[k])
                if tally[best] / len(text_cells) >= VALUE_SHARE:
                    col.kind, col.reason = best, "values"

    # Name parts: consecutive part columns form one person; a repeated part starts a new group.
    group, seen = 0, set()
    for c in sorted(columns):
        col = columns[c]
        if col.part is None:
            continue
        if col.part in seen or not seen:
            group += 1
            seen = set()
        seen.add(col.part)
        col.group = group

    # Review items for classified columns.
    for c, col in columns.items():
        if col.kind is None or col.part is not None:
            continue
        for r in data_rows:
            v = sh.get(r, c)
            if not _replaceable(v, col):
                continue
            s = _surface(v)
            items.append(
                TableItem(
                    kind=col.kind,
                    text=s,
                    sheet=sh.name,
                    cell=f"{col.letter}{r}",
                    context=_context(col.header, _row_context(sh, r, c)),
                    sources=["column"],
                    end=len(s),
                )
            )
    for g in sorted({col.group for col in columns.values() if col.part}):
        cols = [col for col in columns.values() if col.part and col.group == g]
        for r in data_rows:
            parts = _row_parts(sh, r, cols)
            if not parts:
                continue
            name = _person_from_parts({p: v for p, (_, v) in parts.items()})
            first_col = min(cc for cc, _ in parts.values())
            items.append(
                TableItem(
                    kind=Kind.PERSON,
                    text=name,
                    sheet=sh.name,
                    cell=f"{_col_letter(first_col)}{r}",
                    context=_context(
                        " + ".join(col.header for col in cols),
                        _row_context(sh, r, first_col),
                    ),
                    sources=["column"],
                    end=len(name),
                )
            )

    # Free text: rows above the header and data cells of unclassified columns.
    pre_header = range(1, header_rows[0]) if header_rows else range(0)
    for r in [*pre_header, *data_rows]:
        for c in range(1, sh.max_col + 1):
            col = columns[c]
            if r >= data_start and col.kind is not None:
                continue
            v = sh.get(r, c)
            if _cell_type(v) != "text":
                continue
            spans = detect(v)
            cell_texts[(r, c)] = v
            if not spans:
                continue
            findings[(r, c)] = spans
            header = col.header if r >= data_start else ""
            for s in spans:
                items.append(
                    TableItem(
                        kind=s.kind,
                        text=s.text,
                        sheet=sh.name,
                        cell=f"{col.letter}{r}",
                        context=_context(header, _window(v, s.start, s.end)),
                        sources=list(s.sources),
                        start=s.start,
                        end=s.end,
                        normal=s.normal,
                    )
                )

    return SheetScan(
        name=sh.name,
        header_row=header_rows[0] if header_rows else None,
        header_rows=header_rows,
        data_start=data_start,
        columns=columns,
        findings=findings,
        cell_texts=cell_texts,
        items=items,
    )


def _replaceable(v: Any, col: ColumnInfo) -> bool:
    t = _cell_type(v)
    if t == "text":
        return not _is_placeholder(v)
    return t == "number" and col.kind in IDENTIFIER_KINDS


def _row_parts(
    sh: _Sheet, r: int, cols: list[ColumnInfo]
) -> dict[str, tuple[int, str]]:
    out: dict[str, tuple[int, str]] = {}
    for col in cols:
        v = sh.get(r, col.index)
        if _cell_type(v) == "text" and not _is_placeholder(v):
            out[col.part] = (col.index, _surface(v))
    return out


def _open(path: Path) -> tuple[str, list[_Sheet], Any, str, str]:
    """(format, sheets, handle to save, delimiter, encoding)."""
    suffix = path.suffix.lower()
    if suffix in XLSX_SUFFIXES:
        wb = _load_xlsx(path)
        return "xlsx", _xlsx_sheets(wb), wb, "", ""
    if suffix in CSV_SUFFIXES:
        rows, delim, enc = _read_csv(path)
        return "csv", [_csv_sheet(path.stem, rows)], rows, delim, enc
    raise ValueError(f"{path.name}: unsupported table format {suffix!r}")


def scan_table(path: Path) -> TableScan:
    """Classify the columns of every sheet and collect findings for review."""
    path = Path(path)
    fmt, sheets, _, delim, enc = _open(path)
    find_all = _detector()
    return TableScan(
        path=path,
        format=fmt,
        sheets=[_scan_sheet(sh, find_all) for sh in sheets],
        delimiter=delim or ",",
        encoding=enc or "utf-8",
    )


# ---- applying ------------------------------------------------------------


def _table_language(scan: TableScan) -> str:
    text = scan.headers_text() or " ".join(it.text for it in scan.items())
    d = _lang.detect(text)
    if d in ("ru", "en"):
        return d
    cyr = len(re.findall(r"[А-Яа-яЁё]", text))
    lat = len(re.findall(r"[A-Za-z]", text))
    return "ru" if cyr >= lat else "en"


def person_entity(vault: Vault, surface: str) -> Entity:
    """PERSON entity for a name, reusing a known person with the same identity."""
    e = vault.find(Kind.PERSON, surface)
    if e is None:
        matches = [
            x
            for x in vault.entities
            if x.kind == Kind.PERSON and names.same_person(surface, x.canonical)
        ]
        exact = [
            x for x in matches if names.name_key(x.canonical) == names.name_key(surface)
        ]
        if exact:
            e = exact[0]
        elif len(matches) == 1:
            e = matches[0]
    if e is None:
        return vault.get_or_create(Kind.PERSON, surface, canonical=surface)
    e.variants.add(fold(surface))
    return e


def _entity(vault: Vault, kind: Kind, surface: str) -> Entity:
    if kind == Kind.PERSON:
        return person_entity(vault, surface)
    return vault.get_or_create(kind, surface=surface)


def _apply_columns(sh: _Sheet, ss: SheetScan, vault: Vault, lang: str) -> None:
    data_rows = range(ss.data_start, sh.max_row + 1)
    for col in ss.columns.values():
        if col.kind is None or col.part is not None:
            continue
        for r in data_rows:
            v = sh.get(r, col.index)
            if not _replaceable(v, col):
                continue
            s = _surface(v)
            if vault.is_ignored(s):
                continue
            sh.set(r, col.index, vault.label(_entity(vault, col.kind, s), lang))

    for g in sorted({c.group for c in ss.columns.values() if c.part}):
        cols = [c for c in ss.columns.values() if c.part and c.group == g]
        for r in data_rows:
            parts = _row_parts(sh, r, cols)
            if not parts:
                continue
            name = _person_from_parts({p: v for p, (_, v) in parts.items()})
            if vault.is_ignored(name):
                continue
            label = vault.label(person_entity(vault, name), lang)
            for cc, _ in parts.values():
                sh.set(r, cc, label)


def _apply_free_text(
    sh: _Sheet,
    ss: SheetScan,
    vault: Vault,
    lang: str,
    replace: Callable[[str, list[Merged], Vault, str], str],
) -> None:
    pre_header = range(1, ss.header_rows[0]) if ss.header_rows else range(0)
    for r in [*pre_header, *range(ss.data_start, sh.max_row + 1)]:
        for c in range(1, sh.max_col + 1):
            col = ss.columns.get(c)
            if r >= ss.data_start and col is not None and col.kind is not None:
                continue
            v = sh.get(r, c)
            if _cell_type(v) != "text":
                continue
            if ss.cell_texts.get((r, c)) == v:
                spans = ss.findings.get((r, c), [])
            else:
                spans = list(_detector()(v))  # the file changed since the scan
            spans = [s for s in spans if not vault.is_ignored(s.text)]
            new = replace(v, spans, vault, lang)
            if new != v:
                sh.set(r, c, new)


def apply_table(
    path: Path,
    scan: TableScan,
    vault: Vault,
    out_path: Path,
    lang: str | None = None,
) -> None:
    """Write an anonymized copy of the table to out_path."""
    path, out_path = Path(path), Path(out_path)
    lang = lang or _table_language(scan)
    fmt, sheets, handle, delim, _ = _open(path)
    replace = _replacer()
    by_name = {s.name: s for s in scan.sheets}
    pairs = [(sh, by_name[sh.name]) for sh in sheets if sh.name in by_name]
    if fmt == "csv" and not pairs and sheets and scan.sheets:
        pairs = [(sheets[0], scan.sheets[0])]  # renamed csv: one sheet either way

    # Classified columns first, so free-text cells can reuse the people they define.
    for sh, ss in pairs:
        _apply_columns(sh, ss, vault, lang)
    for sh, ss in pairs:
        _apply_free_text(sh, ss, vault, lang, replace)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "xlsx":
        handle.save(out_path)
    else:
        buf = io.StringIO()
        csv.writer(
            buf, delimiter=delim or scan.delimiter, lineterminator="\r\n"
        ).writerows(handle)
        out_path.write_text(buf.getvalue(), encoding="utf-8-sig", newline="")
