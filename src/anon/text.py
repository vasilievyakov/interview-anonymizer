"""Reading and writing plain-text inputs: .md, .txt and .docx."""

from __future__ import annotations

from pathlib import Path

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".text"}
DOCX_SUFFIXES = {".docx"}


def _decode(raw: bytes) -> str:
    # utf-8-sig accepts plain utf-8 too and drops a BOM when present.
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251")


def _table_lines(table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        for cell in row.cells:
            # Merged cells repeat the same underlying element across the row.
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            for p in cell.paragraphs:
                if p.text.strip():
                    lines.append(p.text)
            for inner in cell.tables:
                lines.extend(_table_lines(inner))
    return lines


def _read_docx(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    blocks: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text
            if t.strip():
                blocks.append(t)
        elif tag == "tbl":
            lines = _table_lines(Table(child, doc))
            if lines:
                blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def read_text(path: str | Path) -> str:
    """Text of a .md/.txt (utf-8, utf-8 with BOM, cp1251) or .docx file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in DOCX_SUFFIXES:
        return _read_docx(path)
    if suffix in TEXT_SUFFIXES:
        return _decode(path.read_bytes())
    raise ValueError(f"{path}: unsupported text format {suffix or '(none)'}")


def write_text(path: str | Path, s: str) -> Path:
    path = Path(path)
    path.write_text(s, encoding="utf-8")
    return path


def out_path_for(path: str | Path, suffix: str = "anon") -> Path:
    """interview.md -> interview.anon.md; interview.docx -> interview.anon.md."""
    path = Path(path)
    ext = ".md" if path.suffix.lower() in DOCX_SUFFIXES else path.suffix
    return path.with_name(f"{path.stem}.{suffix}{ext}")
