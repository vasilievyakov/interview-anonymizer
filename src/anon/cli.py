"""Command line: anon scan | review | apply | restore."""

from __future__ import annotations

import argparse
import getpass
import json
import os
import re
import sys
from pathlib import Path

from .agent import quiet
from .i18n import t
from .model import Kind
from .review import (
    REVIEW_FILE,
    Collector,
    Item,
    apply_decisions,
    context_of,
    interactive,
    read_csv,
    write_csv,
)
from .vault import VAULT_DIR, Vault, VaultError

TEXT_EXT = {".md", ".txt", ".docx"}
TABLE_EXT = {".xlsx", ".csv"}
SCAN_STATE = "scan.json"
_LINE = re.compile(r"(?:line|строка) (\d+):")  # line number at the start of a restore problem


def _project(args: argparse.Namespace) -> Path:
    return Path(args.project).resolve()


def _password(project: Path, confirm_new: bool = False) -> str:
    env = os.environ.get("ANON_PASSWORD")
    if env:
        return env
    if confirm_new and not Vault.exists(project):
        first = getpass.getpass(t("password_new"))
        if getpass.getpass(t("password_repeat")) != first:
            sys.exit(t("password_mismatch"))
        return first
    return getpass.getpass(t("password_prompt"))


def _load_vault(project: Path, confirm_new: bool = False) -> tuple[Vault, str]:
    if not Vault.exists(project) and not confirm_new:
        return Vault(), ""
    password = _password(project, confirm_new)
    try:
        return Vault.load(project, password), password
    except VaultError:
        sys.exit(t("wrong_password"))


def _check_files(paths: list[str]) -> list[Path]:
    out = []
    for p in paths:
        path = Path(p).resolve()
        if not path.exists():
            sys.exit(t("file_missing", path=p))
        if path.suffix.lower() not in TEXT_EXT | TABLE_EXT:
            sys.exit(t("unsupported", path=p))
        out.append(path)
    return out


def _out_path(path: Path) -> Path:
    from .text import out_path_for

    if path.suffix.lower() in TABLE_EXT:
        return path.with_name(f"{path.stem}.anon{path.suffix}")
    return out_path_for(path)


# ---- scan ---------------------------------------------------------------


def _say_scanning(path: Path, lang: str, i: int, total: int) -> None:
    # File names often carry the respondent's name: an agent sees only the position.
    if quiet():
        print(t("scanning_quiet", i=i, total=total, lang=lang))
    else:
        print(t("scanning", path=path.name, lang=lang))


def _scan_text(path: Path, collector: Collector, i: int = 1, total: int = 1) -> int:
    from .detect import find_all
    from .lang import languages
    from .text import read_text

    text = read_text(path)
    _say_scanning(path, "+".join(languages(text)), i, total)
    spans = find_all(text)
    for s in spans:
        collector.add(
            s.kind,
            s.text,
            path.name,
            s.sources,
            context_of(text, s.start, s.end),
            s.normal,
        )
    return len(spans)


def _scan_table(path: Path, collector: Collector, i: int = 1, total: int = 1) -> int:
    from .table import scan_table

    _say_scanning(path, "table", i, total)
    scan = scan_table(path)
    n = 0
    for rec in scan.items():
        collector.add(
            Kind(rec.kind),
            rec.text,
            path.name,
            list(getattr(rec, "sources", []) or ["column"]),
            getattr(rec, "context", ""),
            getattr(rec, "normal", ""),
        )
        n += 1
    return n


def _warn_file_names(files: list[Path]) -> None:
    """File names often carry the respondent's name; they travel to the model with the file."""
    from .detect import find_all

    sensitive = {Kind.PERSON, Kind.ORG, Kind.PHONE, Kind.EMAIL}
    hide = quiet()
    for n, path in enumerate(files, 1):
        stem = path.stem.replace("_", " ")
        kinds = sorted({s.kind.value for s in find_all(stem) if s.kind in sensitive})
        if kinds and hide:
            print(
                t(
                    "name_in_filename_quiet",
                    i=n,
                    total=len(files),
                    kinds="/".join(kinds),
                    ext=path.suffix,
                )
            )
        elif kinds:
            print(
                t(
                    "name_in_filename",
                    path=path.name,
                    kinds="/".join(kinds),
                    ext=path.suffix,
                )
            )


def cmd_scan(args: argparse.Namespace) -> None:
    project = _project(args)
    files = _check_files(args.files)
    vault, _ = _load_vault(project)
    collector = Collector()
    total = 0
    for n, path in enumerate(files, 1):
        total += (
            _scan_table(path, collector, n, len(files))
            if path.suffix.lower() in TABLE_EXT
            else _scan_text(path, collector, n, len(files))
        )
    _warn_file_names(files)
    items = collector.items()
    from .keep import load as load_keep

    keep_list = load_keep(project)
    kept_public = 0
    for item in items:
        if all(vault.is_ignored(v) for v in item.variants | {item.text}):
            item.keep = False
        # A finding read as a person is kept only when it is exactly a public name:
        # "Алиса пользуется Ритмо" must stay under review even though "Ритмо" is kept.
        elif any(
            keep_list.matches(v, whole=item.kind == Kind.PERSON)
            for v in item.variants | {item.text}
        ):
            item.keep = False
            item.sources.add("keep")
            kept_public += 1
    review = project / REVIEW_FILE
    write_csv(items, review)
    state = project / VAULT_DIR / SCAN_STATE
    state.parent.mkdir(parents=True, exist_ok=True)
    state.write_text(
        json.dumps({"files": [str(p) for p in files]}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(t("scan_done", n=len(items), files=len(files), review=review.name))
    if kept_public:
        print(t("scan_kept_public", n=kept_public))
    print(t("scan_next_quiet" if quiet() else "scan_next", review=review.name))


# ---- review -------------------------------------------------------------


def cmd_review(args: argparse.Namespace) -> None:
    project = _project(args)
    review = project / REVIEW_FILE
    if not review.exists():
        sys.exit(t("no_scan"))
    items = read_csv(review)
    if args.add:
        kind = Kind(args.kind.upper())
        items.append(
            Item(
                kind=kind,
                text=args.add,
                keep=True,
                variants={args.add},
                sources={"user"},
            )
        )
    elif args.web:
        # The page shows the values; the terminal gets only the URL and the counts.
        from .webreview import serve

        items = serve(project, items, open_browser=not args.no_browser)
    else:
        if quiet():
            sys.exit(t("review_quiet_refused"))
        items = interactive(items)
    write_csv(items, review)
    keep = sum(1 for i in items if i.keep)
    added = sum(1 for i in items if "user" in i.sources)
    print(t("review_done", keep=keep, skip=len(items) - keep, added=added))


def cmd_keep(args: argparse.Namespace) -> None:
    from .keep import add_values, read_values

    project = _project(args)
    if args.values:
        added = add_values(project, args.values)
        print(t("keep_added", n=len(added)))
    values = read_values(project)
    print(t("keep_list", n=len(values)))
    for v in values:
        print(f"  {v}")


def cmd_install_skill(args: argparse.Namespace) -> None:
    """Copy the Claude Code skill and its guard hook out of the package; print the hook snippet.

    The hook runs on the Python of this tool (sys.executable): a clean Mac has no
    working python3 until the developer tools are installed.
    """
    import json
    import os
    import shutil
    from importlib import resources

    base = Path(os.environ.get("CLAUDE_SKILLS_DIR") or Path.home() / ".claude" / "skills")
    dest = base / "anonymize"
    (dest / "hooks").mkdir(parents=True, exist_ok=True)
    src = resources.files("anon") / "claude"
    with resources.as_file(src / "SKILL.md") as f:
        shutil.copyfile(f, dest / "SKILL.md")
    guard = dest / "hooks" / "guard.py"
    with resources.as_file(src / "guard.py") as f:
        shutil.copyfile(f, guard)
    guard.chmod(0o755)
    command = f'"{sys.executable}" "{guard}"'
    snippet = {
        "hooks": {
            "PreToolUse": [
                {
                    "matcher": "Read|Grep|Glob|Edit|MultiEdit|Write|NotebookEdit|Bash",
                    "hooks": [{"type": "command", "command": command}],
                }
            ]
        }
    }
    print(t("skill_installed", dest=dest))
    print()
    print(t("skill_snippet"))
    print()
    print(json.dumps(snippet, ensure_ascii=False, indent=2))


def cmd_list(args: argparse.Namespace) -> None:
    project = _project(args)
    review = project / REVIEW_FILE
    if not review.exists():
        sys.exit(t("no_scan"))
    items = read_csv(review)
    if args.kind:
        items = [i for i in items if i.kind.value == args.kind.upper()]
    order = [Kind.PERSON, Kind.ORG, Kind.LOCATION]
    items.sort(
        key=lambda i: (
            order.index(i.kind) if i.kind in order else len(order),
            i.kind.value,
            -i.count,
            i.text,
        )
    )
    if quiet():
        print(t("values_hidden", review=review.name))
        kinds: dict[str, list[Item]] = {}
        for i in items:
            kinds.setdefault(i.kind.value, []).append(i)
        for kind, group in kinds.items():
            keep = sum(1 for i in group if i.keep)
            print(
                t(
                    "list_kind_quiet",
                    kind=kind,
                    n=len(group),
                    keep=keep,
                    skip=len(group) - keep,
                )
            )
        return
    width = min(36, max((len(i.text) for i in items), default=10))
    for i in items[: args.limit or None]:
        mark = " " if i.keep else "-"
        print(
            f"{mark} {i.kind.value:<9} {i.text[:width]:<{width}}  x{i.count:<3} {','.join(sorted(i.sources))}"
        )


# ---- apply --------------------------------------------------------------


def _apply_text(path: Path, vault: Vault) -> tuple[str, int]:
    from .detect import find_all
    from .lang import label_language
    from .replace import apply_with_report
    from .text import read_text

    text = read_text(path)
    out, report = apply_with_report(text, find_all(text), vault, label_language(text))
    return out, len(report)


def _apply_table(path: Path, vault: Vault, dst: Path) -> set[str]:
    """Writes the anonymized table to dst; returns the folded header cells (column names stay)."""
    from .table import apply_table, scan_table
    from .vault import fold

    scan = scan_table(path)
    apply_table(path, scan, vault, dst)
    return {
        fold(c.header) for sh in scan.sheets for c in sh.columns.values() if c.header
    }


def _count_labels(text: str) -> int:
    import re

    return len(re.findall(r"\[[A-ZА-ЯЁ]+_\d+\]", text))


def _table_text(path: Path) -> str:
    """All cell values of a table as text, for the leak check."""
    if path.suffix.lower() == ".csv":
        return path.read_text(encoding="utf-8-sig")
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def _leak_kinds(leaks: list, items: list[Item], vault: Vault) -> str:
    """Kinds of the leaked values as "PERSON x2, ORG x1", without the values."""
    from collections import Counter

    from .vault import fold

    kind_of: dict[str, str] = {}
    for e in vault.entities:
        for v in e.variants | {e.canonical}:
            kind_of.setdefault(fold(v), e.kind.value)
    for i in items:
        for v in i.variants | {i.text}:
            kind_of.setdefault(fold(v), i.kind.value)
    counts = Counter(kind_of.get(fold(leak.value), "?") for leak in leaks)
    return ", ".join(f"{k} x{n}" for k, n in counts.most_common())


def cmd_apply(args: argparse.Namespace) -> None:
    from .leak import find_leaks
    from .text import write_text

    project = _project(args)
    state = project / VAULT_DIR / SCAN_STATE
    review = project / REVIEW_FILE
    if not state.exists() or not review.exists():
        sys.exit(t("no_scan"))
    files = [Path(p) for p in json.loads(state.read_text(encoding="utf-8"))["files"]]
    vault, password = _load_vault(project, confirm_new=True)
    items = read_csv(review)
    apply_decisions(items, vault)
    # Every surface the human kept must disappear from every output file, even where a
    # detector misses it. For this run they are replaced literally like hand-added values;
    # only the human's own additions stay in the vault for later runs.
    must_hide = [v for i in items if i.keep for v in i.variants | {i.text}]
    persistent_added = list(vault.added)
    known = {(s, k) for s, k in vault.added}
    vault.added.extend(
        (v, i.kind)
        for i in items
        if i.keep
        for v in i.variants | {i.text}
        if (v, i.kind) not in known
    )

    failed = False
    hide = quiet()
    for i, path in enumerate(files, 1):
        dst = _out_path(path)
        # Under an agent a file name may carry a respondent's name: number it instead.
        name = t("file_n", i=i, total=len(files)) if hide else path.name
        if path.suffix.lower() in TABLE_EXT:
            tmp = dst.with_name(dst.name + ".tmp" + dst.suffix)
            headers = _apply_table(path, vault, tmp)
            table_text = _table_text(tmp)
            n = _count_labels(table_text)
            leaks = find_leaks(table_text, vault, must_hide, skip=headers)
            if leaks:
                tmp.unlink(missing_ok=True)
            else:
                tmp.replace(dst)
        else:
            out, _ = _apply_text(path, vault)
            n = _count_labels(out)
            leaks = find_leaks(out, vault, must_hide)
            if not leaks:
                write_text(dst, out)
        if leaks:
            failed = True
            print(t("leak_found", path=name, n=len(leaks)))
            if hide:
                print(
                    t(
                        "leak_lines_quiet",
                        lines=", ".join(
                            str(x) for x in sorted({k.line for k in leaks})
                        ),
                        kinds=_leak_kinds(leaks, items, vault),
                    )
                )
                print(t("leak_hint_quiet", review=review.name))
                continue
            for leak in leaks[:20]:
                print(
                    t(
                        "leak_line",
                        line=leak.line,
                        value=leak.value,
                        context=leak.context,
                    )
                )
            print(t("leak_hint", review=review.name))
        else:
            if hide:
                print(t("apply_file_quiet", src=name, ext="".join(dst.suffixes), n=n))
            else:
                print(t("apply_file", src=path.name, dst=dst.name, n=n))

    vault.added = persistent_added
    vault_path = vault.save(project, password)
    if failed:
        sys.exit(2)
    review.unlink(missing_ok=True)
    state.unlink(missing_ok=True)
    print(t("apply_done", vault=vault_path.relative_to(project)))


# ---- restore ------------------------------------------------------------


def cmd_restore(args: argparse.Namespace) -> None:
    from .restore import restore
    from .text import read_text, write_text

    project = _project(args)
    vault, _ = _load_vault(project, confirm_new=False)
    if not vault.entities and not Vault.exists(project):
        sys.exit(t("no_scan"))
    src = Path(args.file).resolve()
    if not src.exists():
        sys.exit(t("file_missing", path=args.file))
    text = read_text(src)
    restored, problems = restore(text, vault)
    dst = (
        Path(args.output)
        if args.output
        else src.with_name(
            f"{src.stem}.restored{'.md' if src.suffix == '.docx' else src.suffix}"
        )
    )
    write_text(dst, restored)
    n = sum(text.count(label) for label in vault.labels())
    print(t("restore_done", dst=dst.name, n=n))
    if problems and quiet():
        # A problem line quotes the restored value: an agent gets line numbers only.
        lines = sorted({int(m.group(1)) for p in problems if (m := _LINE.match(p))})
        print(
            t(
                "restore_problems_quiet",
                n=len(problems),
                lines=", ".join(map(str, lines)) or "?",
                dst=dst.name,
            )
        )
    elif problems:
        print(t("restore_problems", n=len(problems)))
        for p in problems:
            print(f"    {p}")


def main(argv: list[str] | None = None) -> None:
    # Under an agent or in a pipe stdout is block-buffered: the review page URL and
    # progress lines would only show up after the command ends.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(line_buffering=True)
    parser = argparse.ArgumentParser(
        prog="anon",
        description="Local anonymizer for tables and texts before sending them to an LLM.",
    )
    parser.add_argument(
        "--project",
        default=".",
        help="project folder with the vault (.anon/), default: current folder",
    )
    from . import __version__

    parser.add_argument(
        "--version", action="version", version=f"%(prog)s {__version__}"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("scan", help="find sensitive values and write review.csv")
    p.add_argument("files", nargs="+")
    p.set_defaults(func=cmd_scan)

    p = sub.add_parser("review", help="go through the findings in the terminal")
    p.add_argument("--add", help="add a value the detectors missed")
    p.add_argument(
        "--kind",
        default="PERSON",
        choices=[k.value for k in Kind],
        help="kind for --add",
    )
    p.add_argument(
        "--web",
        action="store_true",
        help="review in a local browser page (127.0.0.1, one-time token)",
    )
    p.add_argument(
        "--no-browser",
        action="store_true",
        help="with --web: only print the URL, do not open the browser",
    )
    p.set_defaults(func=cmd_review)

    p = sub.add_parser(
        "keep",
        help="public names to leave in the text (brands being researched); no values: show the list",
    )
    p.add_argument("values", nargs="*")
    p.set_defaults(func=cmd_keep)

    p = sub.add_parser(
        "install-skill",
        help="install the Claude Code skill and guard hook to ~/.claude/skills/anonymize",
    )
    p.set_defaults(func=cmd_install_skill)

    p = sub.add_parser("list", help="print the review list as a table")
    p.add_argument("--kind", help="only this kind, e.g. PERSON")
    p.add_argument("--limit", type=int, default=0)
    p.set_defaults(func=cmd_list)

    p = sub.add_parser(
        "apply",
        help="replace values with labels, run the leak check, write *.anon.* files",
    )
    p.set_defaults(func=cmd_apply)

    p = sub.add_parser("restore", help="put real values back into the model's answer")
    p.add_argument("file")
    p.add_argument("-o", "--output")
    p.set_defaults(func=cmd_restore)

    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
